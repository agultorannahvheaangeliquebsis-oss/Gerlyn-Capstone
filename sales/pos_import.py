"""Parsing, preview validation, and atomic commit for ZANK POS exports."""

import csv
import io
import re
from collections import defaultdict
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from django.db import transaction
from django.utils import timezone
from django.utils.dateparse import parse_date, parse_datetime
from openpyxl import load_workbook

from billing.models import BillingRecord
from inventory.models import Inventory, Product

from .models import PosSalesImport, SalesItem, SalesTransaction


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MONEY = Decimal('0.01')
ZERO = Decimal('0.00')


class PosImportError(Exception):
    pass


HEADER_ALIASES = {
    'invoice': {
        'invoice', 'invoiceno', 'invoicenumber', 'invoiceornumber',
        'invoiceorno', 'ornumber', 'orno', 'or', 'receiptno',
    },
    'transaction_at': {'transactiondatetime', 'transactiondate', 'datetime', 'date', 'trandate'},
    'terminal': {'terminal', 'terminalid', 'station', 'stationid', 'posid'},
    'cashier': {'cashier', 'cashierid', 'staff', 'staffid', 'userid'},
    'sku': {'sku', 'skucode', 'productsku', 'itemcode'},
    'barcode': {'barcode', 'barcodeno', 'ean', 'upc'},
    'description': {'description', 'productdescription', 'productname', 'itemdescription', 'itemname'},
    'quantity': {'quantity', 'qty', 'quantitysold'},
    'unit_price': {'unitprice', 'price', 'sellingprice'},
    'gross': {'gross', 'grosssales', 'grossamount'},
    'net': {'net', 'netsales', 'netamount', 'amount'},
    'vatable': {'vatable', 'vatablesales', 'vatableamount'},
    'vat_exempt': {'vatexempt', 'vatexemptsales', 'vatexemptamount', 'exemptsales'},
    'zero_rated': {'zerorated', 'zeroratedsales', 'zeroratedamount'},
    'senior_discount': {'seniordiscount', 'seniordiscountamount', 'scdiscount'},
    'pwd_discount': {'pwddiscount', 'pwddiscountamount'},
    'wholesale_discount': {'wholesalediscount', 'wholesalediscountamount'},
    'promo_discount': {'promodiscount', 'promodiscountamount', 'promotiondiscount'},
    'payment': {'payment', 'paymentmethod', 'modeofpayment', 'tender'},
}

PAYMENT_MAP = {
    'cash': SalesTransaction.ModeOfPayment.CASH,
    'gcash': SalesTransaction.ModeOfPayment.GCASH,
    'maya': SalesTransaction.ModeOfPayment.MAYA,
    'paymaya': SalesTransaction.ModeOfPayment.MAYA,
    'qrph': SalesTransaction.ModeOfPayment.QR_PH,
    'card': SalesTransaction.ModeOfPayment.CARD,
    'creditcard': SalesTransaction.ModeOfPayment.CARD,
    'debitcard': SalesTransaction.ModeOfPayment.CARD,
    'chargesale': SalesTransaction.ModeOfPayment.CHARGE_SALE,
    'charge': SalesTransaction.ModeOfPayment.CHARGE_SALE,
}


def _key(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').strip().lower())


def _code_key(value):
    """Case-insensitive product/invoice comparison without erasing punctuation."""
    return str(value or '').strip().casefold()


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _money(value, label, errors, *, required=False):
    raw = _text(value).replace(',', '').replace('₱', '')
    if not raw:
        if required:
            errors.append(f'{label} is required.')
        return None
    try:
        number = Decimal(raw)
    except InvalidOperation:
        errors.append(f'{label} must be a number.')
        return None
    if not number.is_finite() or number < 0:
        errors.append(f'{label} must be zero or greater.')
        return None
    return number.quantize(MONEY, rounding=ROUND_HALF_UP)


def _transaction_datetime(value, errors):
    result = None
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, date):
        result = datetime.combine(value, time.min)
    else:
        raw = _text(value)
        result = parse_datetime(raw)
        if result is None:
            parsed_date = parse_date(raw)
            if parsed_date:
                result = datetime.combine(parsed_date, time.min)
        if result is None:
            for fmt in (
                '%m/%d/%Y %H:%M:%S', '%m/%d/%Y %H:%M',
                '%m/%d/%Y %I:%M %p', '%d/%m/%Y %H:%M:%S',
                '%d/%m/%Y %H:%M', '%Y-%m-%d %H:%M:%S',
            ):
                try:
                    result = datetime.strptime(raw, fmt)
                    break
                except ValueError:
                    pass
    if result is None:
        errors.append('Transaction date/time is missing or invalid.')
        return None
    if timezone.is_naive(result):
        result = timezone.make_aware(result, timezone.get_current_timezone())
    return result


def _sheet_rows(upload, extension):
    upload.seek(0)
    if extension == '.csv':
        raw = upload.read()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise PosImportError('The CSV must be UTF-8 encoded.') from exc
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        try:
            return list(csv.reader(io.StringIO(text), dialect))
        except csv.Error as exc:
            raise PosImportError(f'Could not read the CSV: {exc}.') from exc
    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    except Exception as exc:
        raise PosImportError('Could not read that Excel workbook. Confirm it is a valid .xlsx file.') from exc


def parse_pos_file(upload):
    """Return JSON-serializable normalized rows; structural errors raise."""
    filename = upload.name or ''
    extension = '.' + filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if extension not in {'.csv', '.xlsx'}:
        raise PosImportError('Choose a CSV or Excel (.xlsx) file.')
    if upload.size > MAX_UPLOAD_BYTES:
        raise PosImportError('The POS file cannot exceed 10 MB.')
    rows = _sheet_rows(upload, extension)
    while rows and not any(_text(cell) for cell in rows[0]):
        rows.pop(0)
    if not rows:
        raise PosImportError('The POS file is empty.')

    alias_lookup = {alias: canonical for canonical, aliases in HEADER_ALIASES.items() for alias in aliases}
    columns = {}
    for index, header in enumerate(rows[0]):
        canonical = alias_lookup.get(_key(header))
        if canonical and canonical not in columns:
            columns[canonical] = index
    required = ['invoice', 'transaction_at', 'quantity', 'unit_price', 'payment']
    missing = [name for name in required if name not in columns]
    if 'sku' not in columns and 'barcode' not in columns:
        missing.append('SKU or barcode')
    if missing:
        raise PosImportError('Missing required column(s): ' + ', '.join(missing) + '.')

    normalized = []
    for row_number, values in enumerate(rows[1:], start=2):
        if not any(_text(cell) for cell in values):
            continue

        def cell(name):
            index = columns.get(name)
            return values[index] if index is not None and index < len(values) else None

        errors = []
        invoice = _text(cell('invoice'))
        if not invoice:
            errors.append('Invoice/OR number is required.')
        elif len(invoice) > 100:
            errors.append('Invoice/OR number cannot exceed 100 characters.')
        occurred = _transaction_datetime(cell('transaction_at'), errors)
        quantity_raw = _text(cell('quantity')).replace(',', '')
        try:
            quantity_decimal = Decimal(quantity_raw)
            quantity = int(quantity_decimal)
            if quantity_decimal != quantity or quantity <= 0:
                raise ValueError
        except (InvalidOperation, ValueError):
            quantity = None
            errors.append('Quantity must be a positive whole number.')

        money = {}
        for name, label in (
            ('unit_price', 'Unit price'), ('gross', 'Gross sales'), ('net', 'Net sales'),
            ('vatable', 'VATable amount'), ('vat_exempt', 'VAT-exempt amount'),
            ('zero_rated', 'Zero-rated amount'), ('senior_discount', 'Senior discount'),
            ('pwd_discount', 'PWD discount'), ('wholesale_discount', 'Wholesale discount'),
            ('promo_discount', 'Promo discount'),
        ):
            money[name] = _money(cell(name), label, errors, required=name == 'unit_price')

        payment_raw = _text(cell('payment'))
        payment = PAYMENT_MAP.get(_key(payment_raw))
        if not payment:
            errors.append(
                f'Unsupported payment method "{payment_raw or "blank"}". '
                'Use Cash, GCash, Maya, QR Ph, Card, or Charge Sale.'
            )

        sku = _text(cell('sku'))
        barcode = _text(cell('barcode'))
        if not sku and not barcode:
            errors.append('SKU or barcode is required.')
        if len(sku) > 50:
            errors.append('SKU cannot exceed 50 characters.')
        if len(barcode) > 64:
            errors.append('Barcode cannot exceed 64 characters.')

        terminal = _text(cell('terminal'))
        cashier = _text(cell('cashier'))
        if len(terminal) > 100:
            errors.append('Terminal/station ID cannot exceed 100 characters.')
        if len(cashier) > 100:
            errors.append('Cashier/staff ID cannot exceed 100 characters.')
        for name, label in (('unit_price', 'Unit price'), ('gross', 'Gross sales'),
                            ('net', 'Net sales'), ('vatable', 'VATable amount'),
                            ('vat_exempt', 'VAT-exempt amount'), ('zero_rated', 'Zero-rated amount'),
                            ('senior_discount', 'Senior discount'), ('pwd_discount', 'PWD discount'),
                            ('wholesale_discount', 'Wholesale discount'), ('promo_discount', 'Promo discount')):
            value = money.get(name)
            limit = Decimal('99999999.99') if name == 'unit_price' else Decimal('9999999999.99')
            if value is not None and value > limit:
                errors.append(f'{label} is too large.')

        discount = sum((money[name] or ZERO) for name in (
            'senior_discount', 'pwd_discount', 'wholesale_discount', 'promo_discount'
        ))
        gross = money['gross']
        if gross is None and money['unit_price'] is not None and quantity:
            gross = (money['unit_price'] * quantity).quantize(MONEY)
        net = money['net']
        if net is None and gross is not None:
            net = max(ZERO, gross - discount).quantize(MONEY)
        if net is None:
            errors.append('Net sales could not be determined from this row.')
        for value, label in ((gross, 'Gross sales'), (net, 'Net sales')):
            if value is not None and value > Decimal('9999999999.99'):
                errors.append(f'{label} is too large.')

        normalized.append({
            'row_number': row_number, 'invoice': invoice,
            'transaction_at': occurred.isoformat() if occurred else '',
            'terminal': terminal, 'cashier': cashier,
            'sku': sku, 'barcode': barcode, 'description': _text(cell('description')),
            'quantity': quantity, 'payment': payment or '', 'payment_raw': payment_raw,
            **{name: str(value) if value is not None else '' for name, value in money.items()},
            'gross': str(gross) if gross is not None else '',
            'net': str(net) if net is not None else '',
            'errors': errors,
        })
    if not normalized:
        raise PosImportError('The POS file contains no sales rows.')
    if len(normalized) > 50000:
        raise PosImportError('The POS file cannot contain more than 50,000 sales rows.')
    return normalized


def _d(value):
    return Decimal(value or '0.00')


def build_preview(rows, branch):
    """Resolve products, duplicates and stock for a branch."""
    products = list(Product.objects.all())
    by_sku = defaultdict(list)
    by_barcode = defaultdict(list)
    for product in products:
        by_sku[_code_key(product.sku_code)].append(product)
        if product.barcode:
            by_barcode[_code_key(product.barcode)].append(product)

    existing = {
        _code_key(value) for value in SalesTransaction.objects.filter(
            branch=branch, pos_invoice_number__isnull=False,
        ).values_list('pos_invoice_number', flat=True)
    }
    groups = {}
    for source_row in rows:
        row = dict(source_row)
        row['errors'] = list(source_row.get('errors', []))
        candidates = (
            set(by_sku.get(_code_key(row['sku']), []))
            | set(by_barcode.get(_code_key(row['barcode']), []))
        )
        if len(candidates) == 1:
            product = candidates.pop()
            row['product_id'] = product.pk
            row['product_name'] = product.product_name
        elif not candidates:
            row['product_id'] = None
            row['product_name'] = ''
            row['errors'].append(
                f'No Product matches SKU "{row["sku"]}" or barcode "{row["barcode"]}".'
            )
        else:
            row['product_id'] = None
            row['product_name'] = ''
            row['errors'].append('SKU/barcode matches more than one Product; correct the product records first.')
        key = _code_key(row['invoice']) or f'__invalid_row_{row["row_number"]}'
        group = groups.setdefault(key, {
            'invoice': row['invoice'], 'rows': [], 'errors': [], 'duplicate': False,
        })
        group['rows'].append(row)

    stock_needed = defaultdict(int)
    for group in groups.values():
        first = group['rows'][0]
        group['duplicate'] = bool(group['invoice'] and _code_key(group['invoice']) in existing)
        for field, label in (('transaction_at', 'date/time'), ('payment', 'payment method'),
                             ('terminal', 'terminal ID'), ('cashier', 'cashier ID')):
            if any(row[field] != first[field] for row in group['rows'][1:]):
                group['errors'].append(f'Invoice rows have inconsistent {label}.')
        if any(row['errors'] for row in group['rows']):
            group['errors'].append('One or more rows in this invoice are invalid.')
        group['gross'] = sum((_d(row['gross']) for row in group['rows']), ZERO)
        group['net'] = sum((_d(row['net']) for row in group['rows']), ZERO)
        for name in ('vatable', 'vat_exempt', 'zero_rated', 'senior_discount',
                     'pwd_discount', 'wholesale_discount', 'promo_discount'):
            group[name] = sum((_d(row[name]) for row in group['rows']), ZERO)
        if not group['errors'] and not group['duplicate']:
            for row in group['rows']:
                stock_needed[row['product_id']] += row['quantity']

    available = {
        inv.product_id: inv.available_qty
        for inv in Inventory.objects.sellable().filter(branch=branch)
    }
    for product_id, needed in stock_needed.items():
        have = available.get(product_id, 0)
        if needed > have:
            for group in groups.values():
                if group['duplicate']:
                    continue
                affected = [r for r in group['rows'] if r.get('product_id') == product_id]
                if affected:
                    name = affected[0]['product_name']
                    message = f'Not enough stock for {name}: available {have}, file requires {needed}.'
                    group['errors'].append(message)
                    for row in affected:
                        row['errors'].append(message)

    transactions = list(groups.values())
    for group in transactions:
        # Transaction-level problems (for example inconsistent payment
        # methods across one invoice) must appear on the actual preview rows;
        # an "Invalid" pill without its reason would be a silent discard.
        for message in group['errors']:
            for row in group['rows']:
                if message not in row['errors']:
                    row['errors'].append(message)
        group['status'] = 'invalid' if group['errors'] else 'duplicate' if group['duplicate'] else 'valid'
        for row in group['rows']:
            row['status'] = group['status'] if group['status'] != 'valid' else ('invalid' if row['errors'] else 'valid')

    flat_rows = [row for group in transactions for row in group['rows']]
    valid_groups = [group for group in transactions if group['status'] == 'valid']
    return {
        'transactions': transactions,
        'rows': flat_rows,
        'summary': {
            'total_rows': len(flat_rows),
            'valid_rows': sum(len(group['rows']) for group in valid_groups),
            'invalid_rows': sum(1 for row in flat_rows if row['status'] == 'invalid'),
            'duplicates': sum(1 for group in transactions if group['status'] == 'duplicate'),
            'valid_transactions': len(valid_groups),
            'gross_total': sum((group['gross'] for group in valid_groups), ZERO),
            'net_total': sum((group['net'] for group in valid_groups), ZERO),
            'quantity_total': sum(row['quantity'] for group in valid_groups for row in group['rows']),
        },
    }


def commit_pos_import(*, rows, branch, filename, actor):
    """Revalidate and save every eligible invoice as one atomic operation."""
    from .views import _deduct_stock

    with transaction.atomic():
        # Serialize concurrent imports for this branch before duplicate and
        # stock checks. The DB uniqueness constraint is the final guard.
        branch.__class__.objects.select_for_update().get(pk=branch.pk)
        preview = build_preview(rows, branch)
        valid = [group for group in preview['transactions'] if group['status'] == 'valid']
        summary = preview['summary']
        if not valid:
            raise PosImportError('There are no valid, non-duplicate transactions to import.')

        errors = [
            {'row': row['row_number'], 'invoice': row['invoice'], 'errors': row['errors']}
            for row in preview['rows'] if row['status'] == 'invalid'
        ]
        batch = PosSalesImport.objects.create(
            original_filename=filename[:255], branch=branch, imported_by=actor,
            successful_transactions=len(valid), skipped_duplicates=summary['duplicates'],
            failed_rows=summary['invalid_rows'], total_rows=summary['total_rows'],
            gross_total=summary['gross_total'], net_total=summary['net_total'],
            validation_errors=errors,
        )
        for group in valid:
            first = group['rows'][0]
            discount_type = SalesTransaction.DiscountType.NONE
            if group['pwd_discount'] > 0:
                discount_type = SalesTransaction.DiscountType.PWD
            elif group['senior_discount'] > 0:
                discount_type = SalesTransaction.DiscountType.SENIOR
            total_discount = sum((group[name] for name in (
                'senior_discount', 'pwd_discount', 'wholesale_discount', 'promo_discount'
            )), ZERO)
            vatable = group['vatable']
            vat_amount = (vatable * Decimal('0.12') / Decimal('1.12')).quantize(MONEY)
            occurred = datetime.fromisoformat(first['transaction_at'])
            txn = SalesTransaction.objects.create(
                branch=branch, staff=actor if actor.role == 'STAFF' else None,
                recorded_by=actor, transaction_date=occurred,
                transaction_type=SalesTransaction.TransactionType.PHYSICAL,
                fulfillment_type=SalesTransaction.FulfillmentType.PICKUP,
                mode_of_payment=first['payment'],
                order_status=SalesTransaction.OrderStatus.COMPLETED,
                completed_at=occurred, total_amount=group['net'],
                total_discount_amount=total_discount, total_vat_amount=vat_amount,
                source='ZANK POS', pos_invoice_number=group['invoice'],
                pos_terminal_id=first['terminal'], pos_cashier_id=first['cashier'],
                pos_gross_sales=group['gross'], pos_vatable_sales=group['vatable'],
                pos_vat_exempt_sales=group['vat_exempt'],
                pos_zero_rated_sales=group['zero_rated'],
                pos_senior_discount=group['senior_discount'],
                pos_pwd_discount=group['pwd_discount'],
                pos_wholesale_discount=group['wholesale_discount'],
                pos_promo_discount=group['promo_discount'], pos_import=batch,
                discount_type=discount_type,
            )
            for row in group['rows']:
                quantity = row['quantity']
                net = _d(row['net'])
                unit_net = (net / quantity).quantize(MONEY, rounding=ROUND_HALF_UP)
                row_discount = sum((_d(row[name]) for name in (
                    'senior_discount', 'pwd_discount', 'wholesale_discount', 'promo_discount'
                )), ZERO)
                row_vat = (_d(row['vatable']) * Decimal('0.12') / Decimal('1.12')).quantize(MONEY)
                SalesItem.objects.create(
                    transaction=txn, product_id=row['product_id'], quantity_sold=quantity,
                    unit_price=unit_net,
                    is_vat_exempt_line=bool(_d(row['vat_exempt']) or _d(row['zero_rated'])),
                    vat_amount=row_vat, discount_amount=row_discount,
                    pos_sku=row['sku'], pos_barcode=row['barcode'],
                    pos_description=row['description'],
                    pos_unit_price=_d(row['unit_price']),
                    pos_gross_sales=_d(row['gross']), pos_net_sales=net,
                    pos_vatable_sales=_d(row['vatable']),
                    pos_vat_exempt_sales=_d(row['vat_exempt']),
                    pos_zero_rated_sales=_d(row['zero_rated']),
                    pos_senior_discount=_d(row['senior_discount']),
                    pos_pwd_discount=_d(row['pwd_discount']),
                    pos_wholesale_discount=_d(row['wholesale_discount']),
                    pos_promo_discount=_d(row['promo_discount']),
                )
            _deduct_stock(txn, actor)
            billing = txn.billing_record
            billing.amount_due = group['net']
            billing.payment_method = first['payment']
            billing.status = BillingRecord.PaymentStatus.PAID
            billing.confirmed_by = actor
            billing.date_confirmed = timezone.now()
            billing.save()
        return batch, preview
