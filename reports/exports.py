"""
reports/exports.py — the PDF and Excel exports of the four management reports
(Sales, Billing, Inventory, Procurement).

Every export is one ReportDocument: a header block (store, report, branch,
period, who generated it and when), a Summary, and a detail table grouped by
branch with per-branch subtotals and a grand total. The PDF and the Excel
workbook are two renderers over that same object, so the two files cannot
disagree with each other.

Why the Summary cannot disagree with the detail table or the dashboard
---------------------------------------------------------------------
Each builder resolves its row set ONCE, through the same reports.services /
reports.views function the report page itself calls (scoped_sales_in_range,
scoped_billing_in_range, ...), and computes every Summary figure from that
queryset using the same helpers the page's KPI cards use (sales_range_totals,
billing_figures, procurement_figures, inventory_figures). No figure here has its
own definition of "revenue", "low stock" or "procurement cost" — see
reports/services.py and docs/METRICS.md. reports.tests_exports asserts the
match against the dashboard and against the detail rows.

PDF rendering follows the project's other PDFs (sales receipts, purchase
orders): xhtml2pdf, which is pure Python and needs no system libraries. Its
base font has no peso glyph, so PDF money is written "PHP 1,234.00"; the Excel
workbook uses real numbers with a number format.
"""
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

from django.db.models import Count, Q, Sum
from django.db.models.functions import Coalesce
from django.template.loader import render_to_string
from django.utils import timezone

from billing.models import BillingRecord
from branches.models import Branch
from core.people import person_display_name
from core.scoping import branch_ids_for_user
from inventory.alerts import APPROACHING_EXPIRY_DAYS, EXPIRING_SOON_DAYS, _expiry_tier, stock_status
from procurement.models import DeliveryReceipt, Procurement, SupplierReturn, SupplierReturnItem
from reports.services import (
    MONEY_FIELD,
    PIPELINE_ORDER_STATUSES,
    owner_branch_choice,
    procurement_cost_total,
    revenue_sales,
    sales_range_totals,
    scoped_billing_in_range,
    scoped_inventory,
    scoped_procurement_in_range,
    scoped_sales_in_range,
)
from reports.views import (
    _inventory_unit_cost,
    billing_figures,
    inventory_figures,
    nearest_floor_expiry,
    procurement_figures,
)
from sales.models import Refund, SalesTransaction

STORE_NAME = 'Gerlyn Variety Store'
APP_NAME = 'GerlynPLUS'

REPORT_KINDS = ('sales', 'billing', 'inventory', 'procurement')

# xhtml2pdf lays a table out in memory and a gunicorn worker is held for the
# whole render (Procfile: --timeout 60). Measured here: ~19 ms a row, so 1,440
# sales rows took 27 s against 0.8 s for the same rows as Excel. 800 rows is
# ~15 s, which leaves headroom for a slower host. Past it the export refuses and
# points at Excel/CSV, which have no such ceiling; a shorter date range or a
# branch pick brings any report back under it.
MAX_PDF_ROWS = 800

ZERO = Decimal('0.00')

# --- Payment-method grouping -------------------------------------------------
#
# The Sales and Billing summaries answer "how much came in by cash, GCash, Maya
# and card". The payment models carry nine methods, so they are folded into
# those four rows plus "Other methods" — nothing is dropped, which is what lets
# the rows add up to the headline total. Cash on Delivery (cash handed over at
# the door) and Pay at Store (settled at the counter) count as Cash; GCash and
# Maya each keep their own row. The
# grouping is printed on the report itself (see PAYMENT_GROUP_NOTE) so a reader
# never has to guess what a row contains.
CASH = 'Cash'
GCASH = 'GCash'
MAYA = 'Maya'
CARD = 'Card'
OTHER_METHODS = 'Other methods'

PAYMENT_GROUP_ORDER = (CASH, GCASH, MAYA, CARD, OTHER_METHODS)
_CASH_METHODS = frozenset({'Cash', 'Cash on Delivery', 'Pay at Store'})

PAYMENT_GROUP_NOTE = (
    'Payment methods: Cash includes Cash on Delivery and Pay at Store. Other methods = QR Ph, '
    'Charge Sale and Bank Transfer.'
)


def payment_group(method):
    """Which of the summary rows a recorded payment method falls in."""
    if method in _CASH_METHODS:
        return CASH
    if method == 'GCash':
        return GCASH
    if method == 'Maya':
        return MAYA
    if method == 'Card':
        return CARD
    return OTHER_METHODS


# =============================================================================
# The document model
# =============================================================================

@dataclass
class Column:
    key: str
    label: str
    align: str = 'left'     # 'left' | 'right' | 'center'
    width: int = 10         # relative weight; the renderers normalise to 100%


@dataclass
class SummaryBlock:
    """One titled table in the Summary: column headers, then rows of typed
    values (str / int / Decimal / None). The last row is a total when
    `total_row` is set, and is drawn bold."""
    title: str
    headers: list
    rows: list
    total_row: bool = False


@dataclass
class Group:
    """One branch's rows in the detail table. `subtotal` holds the values for
    the columns AFTER the label span (see ReportDocument.label_span); None when
    the report covers a single branch, where it would only repeat the grand
    total."""
    title: str
    rows: list
    subtotal_label: str = ''
    subtotal: list = None


@dataclass
class ReportDocument:
    kind: str
    title: str
    orientation: str                    # 'landscape' | 'portrait' — each report's own choice
    branch_label: str
    period_label: str
    generated_by: str
    generated_at: datetime
    columns: list
    label_span: int                     # leading columns a subtotal / grand-total label spans
    headline: list                      # [(label, value)] — the big figures at the top
    summary: list                       # [SummaryBlock]
    groups: list
    grand_total: list = None            # values for the columns after `label_span`
    grand_total_label: str = 'Grand total'
    notes: list = field(default_factory=list)
    detail_title: str = 'Details'
    sheet_name: str = 'Detail'

    @property
    def row_count(self):
        return sum(len(g.rows) for g in self.groups)

    @property
    def is_empty(self):
        return self.row_count == 0

    @property
    def generated_label(self):
        return f'{self.generated_at:%d %b %Y}, {_clock(self.generated_at)}'


def _clock(moment):
    """12-hour clock without a leading zero ("3:45 PM") — %-I is not portable
    to Windows, where this also runs in development."""
    return f'{moment.hour % 12 or 12}:{moment.minute:02d} {"AM" if moment.hour < 12 else "PM"}'


# =============================================================================
# Value formatting (PDF text). The Excel renderer keeps values typed instead.
# =============================================================================

def fmt_money(value, prefix=True):
    text = f'{Decimal(value):,.2f}'
    return f'PHP {text}' if prefix else text


def fmt_cell(value, money_prefix=False):
    """A typed cell value as the text the PDF prints. Decimals are always money
    in this module; detail tables leave the "PHP" to the column heading, the
    Summary spells it out."""
    if value is None or value == '':
        return '—'
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    if isinstance(value, Decimal):
        return fmt_money(value, prefix=money_prefix)
    if isinstance(value, int):
        return f'{value:,}'
    if isinstance(value, datetime):
        return f'{timezone.localtime(value):%Y-%m-%d %H:%M}'   # as the report pages and CSV write it
    if isinstance(value, date):
        return f'{value:%Y-%m-%d}'
    return str(value)


def _plural(n, word):
    return f'{n:,} {word}{"" if n == 1 else "s"}'


# =============================================================================
# Header: branch and period labels
# =============================================================================

def _branch_name(pk):
    return Branch.objects.filter(pk=pk).values_list('branch_name', flat=True).first()


def branch_label(user, branch_id=None):
    """The branch the report covers, as the header prints it."""
    ids = branch_ids_for_user(user)
    if ids is None:                                  # Owner
        chosen = owner_branch_choice(user, branch_id)
        if not chosen:
            return 'All branches'
        return _branch_name(chosen) or 'Unknown branch'
    # Manager/Staff: their own branch always wins; an unassigned account sees
    # nothing (fails closed) and the header says why.
    return (_branch_name(ids[0]) if ids else None) or 'No branch assigned'


_PRESET_NAMES = {'today': 'Today', 'yesterday': 'Yesterday', '7': 'Last 7 days', '30': 'Last 30 days'}


def period_label(active_range, start, end):
    """The date range as the header prints it: the preset's name when it has one
    plus the exact dates, so a printed page is unambiguous the next day."""
    if start is None:
        return f'All time (through {end:%d %b %Y})'
    dates = f'{end:%d %b %Y}' if start == end else f'{start:%d %b %Y} – {end:%d %b %Y}'
    name = _PRESET_NAMES.get(active_range)
    return f'{name} ({dates})' if name else dates


def _money(value):
    """A money figure as Decimal. The page helpers sum with Python's sum(),
    which yields int 0 for no rows, and an int would print "0" where money
    prints "0.00"."""
    return value if isinstance(value, Decimal) else Decimal(value)


def _document_header(user, branch_id, now):
    return {
        'branch_label': branch_label(user, branch_id),
        'generated_by': person_display_name(user) or user.username,
        'generated_at': timezone.localtime(now or timezone.now()),
    }


def _by_branch(rows, branch_of):
    """Group `rows` by branch preserving their order, the unassigned group last.
    Returns [(branch_or_None, [rows])]."""
    order, buckets = [], {}
    for row in rows:
        branch = branch_of(row)
        key = branch.pk if branch else None
        if key not in buckets:
            buckets[key] = (branch, [])
            order.append(key)
        buckets[key][1].append(row)
    named = sorted((k for k in order if k is not None), key=lambda k: buckets[k][0].branch_name.lower())
    return [buckets[k] for k in named + ([None] if None in buckets else [])]


def _group_title(branch):
    return branch.branch_name if branch else 'No branch assigned'


# =============================================================================
# Sales
# =============================================================================

SALES_COLUMNS = [
    Column('order', 'Order No.', width=11),
    Column('date', 'Date', width=12),
    Column('customer', 'Customer', width=15),
    Column('recorded_by', 'Recorded by', width=11),
    Column('channel', 'Channel', width=7),
    Column('payment', 'Payment', width=10),
    Column('status', 'Status', width=8),
    Column('amount', 'Amount (PHP)', 'right', 9),
    Column('refunded', 'Refunded (PHP)', 'right', 9),
    Column('net', 'Net (PHP)', 'right', 8),
]


def _sales_row(t):
    completed = t.order_status == SalesTransaction.OrderStatus.COMPLETED
    recorder = t.recorded_by or t.staff
    return [
        t.order_number,
        timezone.localtime(t.transaction_date),
        person_display_name(t.customer) or 'Walk-in',      # as the Sales Report page words it
        person_display_name(recorder) or '—',
        'Online' if t.transaction_type == SalesTransaction.TransactionType.ONLINE else 'Walk-in',
        t.mode_of_payment,
        t.order_status,
        t.total_amount,
        t.refunded or None,
        (t.total_amount - t.refunded) if completed else None,
    ]


def _sales_best_branch(rows_qs, completed_qs):
    """(name, net sales) of the branch with the highest Net Sales in scope, or
    None. Same definition as the headline: Completed gross less approved
    refunds, per branch."""
    gross = dict(completed_qs.order_by().values_list('branch_id').annotate(total=Sum('total_amount')))
    refunds = dict(
        Refund.objects.filter(transaction__in=rows_qs, status=Refund.Status.APPROVED)
        .order_by().values_list('transaction__branch_id').annotate(total=Sum('total_amount'))
    )
    net = {bid: total - refunds.get(bid, ZERO) for bid, total in gross.items() if bid is not None}
    if not net:
        return None
    names = Branch.objects.in_bulk(net.keys())
    best = min(net, key=lambda bid: (-net[bid], names[bid].branch_name.lower()))
    return names[best].branch_name, net[best]


def build_sales_document(user, start, end, active_range='30', branch_id=None, now=None):
    # scoped_sales_in_range hands the id straight to Branch.objects.filter(pk=...),
    # which raises ValueError on a junk ?branch=abc; the other builders' row
    # sets clean it themselves (reports.services.owner_branch_choice).
    branch_id = owner_branch_choice(user, branch_id)
    # The one row set the Sales Report page, its JSON mirror and its CSV read.
    rows_qs = scoped_sales_in_range(user, start, end, branch_id=branch_id)
    totals = sales_range_totals(rows_qs)

    # --- Summary, all from rows_qs -------------------------------------------
    by_status = dict(rows_qs.order_by().values_list('order_status').annotate(n=Count('pk')))
    by_channel = dict(rows_qs.order_by().values_list('transaction_type').annotate(n=Count('pk')))
    refunded_orders = (
        Refund.objects.filter(transaction__in=rows_qs, status=Refund.Status.APPROVED)
        .order_by().values('transaction_id').distinct().count()
    )
    completed_qs = revenue_sales(rows_qs)

    payments = {name: [0, ZERO] for name in PAYMENT_GROUP_ORDER}
    for row in completed_qs.order_by().values('mode_of_payment').annotate(n=Count('pk'), amount=Sum('total_amount')):
        bucket = payments[payment_group(row['mode_of_payment'])]
        bucket[0] += row['n']
        bucket[1] += row['amount']

    all_branches = branch_ids_for_user(user) is None and not owner_branch_choice(user, branch_id)
    best = _sales_best_branch(rows_qs, completed_qs) if all_branches else None

    headline = [
        ('Total orders', totals['transactions']),
        ('Gross sales', totals['gross_sales']),
        ('Total refunds', totals['refund_value']),
        ('Net sales', totals['net_sales']),
    ]
    if best:
        headline.append(('Best branch', f'{best[0]} · {fmt_money(best[1])}'))

    in_progress = sum(by_status.get(s, 0) for s in PIPELINE_ORDER_STATUSES)
    summary = [
        SummaryBlock('Orders', ['Measure', 'Orders'], [
            ['Total orders', totals['transactions']],
            ['Completed', by_status.get(SalesTransaction.OrderStatus.COMPLETED, 0)],
            ['Cancelled', by_status.get(SalesTransaction.OrderStatus.CANCELLED, 0)],
            ['Refunded (with an approved refund)', refunded_orders],
            ['In progress (pending, approved, preparing)', in_progress],
            ['Walk-in', by_channel.get(SalesTransaction.TransactionType.PHYSICAL, 0)],
            ['Online', by_channel.get(SalesTransaction.TransactionType.ONLINE, 0)],
        ]),
        SummaryBlock('Sales', ['Measure', 'Amount'], [
            ['Gross sales (completed orders)', totals['gross_sales']],
            ['Total refunds (approved)', totals['refund_value']],
            ['Net sales', totals['net_sales']],
            ['Committed, not yet sales (in progress)', totals['pipeline_value']],
        ], total_row=False),
        SummaryBlock('Payments received (completed orders, before refunds)',
                     ['Method', 'Orders', 'Amount'],
                     [[name, payments[name][0], payments[name][1]] for name in PAYMENT_GROUP_ORDER]
                     + [['Total', totals['orders'], totals['gross_sales']]],
                     total_row=True),
    ]

    # --- Detail: the same rows, grouped by branch -----------------------------
    detail = (
        rows_qs.select_related('recorded_by')
        .annotate(refunded=Coalesce(
            Sum('refunds__total_amount', filter=Q(refunds__status=Refund.Status.APPROVED)),
            ZERO, output_field=MONEY_FIELD,
        ))
        .order_by('branch__branch_name', 'branch_id', '-transaction_date', '-pk')
    )
    buckets = _by_branch(detail, lambda t: t.branch)
    multi = len(buckets) > 1
    groups = []
    for branch, rows in buckets:
        group = Group(title=_group_title(branch), rows=[_sales_row(t) for t in rows])
        if multi:
            scoped = rows_qs.filter(branch=branch) if branch else rows_qs.filter(branch__isnull=True)
            t = sales_range_totals(scoped)
            group.subtotal_label = f'Subtotal · {group.title} · {_plural(t["orders"], "completed order")}'
            group.subtotal = [t['gross_sales'], t['refund_value'], t['net_sales']]
        groups.append(group)

    return ReportDocument(
        kind='sales', title='Sales Report', orientation='landscape',
        period_label=period_label(active_range, start, end),
        columns=SALES_COLUMNS, label_span=7,
        headline=headline, summary=summary, groups=groups,
        grand_total=[totals['gross_sales'], totals['refund_value'], totals['net_sales']],
        grand_total_label=f'Grand total · {_plural(totals["orders"], "completed order")}',
        notes=[
            'Totals count Completed orders only. Amounts shown for Cancelled and in-progress '
            'orders are for reference and are not part of any total.',
            'Net sales = gross sales less approved refunds.',
            PAYMENT_GROUP_NOTE,
        ],
        detail_title='Sales detail', sheet_name='Sales Detail',
        **_document_header(user, branch_id, now),
    )


# =============================================================================
# Billing
# =============================================================================

BILLING_COLUMNS = [
    Column('order', 'Order No.', width=12),
    Column('date', 'Date recorded', width=14),
    Column('method', 'Method', width=11),
    Column('reference', 'Reference No.', width=15),
    Column('status', 'Status', width=11),
    Column('verified_by', 'Verified by', width=13),
    Column('amount', 'Amount due (PHP)', 'right', 12),
    Column('verified', 'Verified (PHP)', 'right', 12),
]

# The page calls a payment that has been marked Paid "Collected"; the summary
# calls it Verified — a billing clerk checks the payment proof, then marks it
# Paid. Same status, same records, same amount (BillingRecord.PaymentStatus.PAID).
_BILLING_STATUS_ORDER = (
    ('Verified (Paid)', [BillingRecord.PaymentStatus.PAID]),
    ('Pending', [BillingRecord.PaymentStatus.PENDING]),
    ('Rejected', [BillingRecord.PaymentStatus.REJECTED]),
    ('Cancelled', [BillingRecord.PaymentStatus.CANCELLED]),
    ('Partially refunded', [BillingRecord.PaymentStatus.PARTIALLY_REFUNDED]),
    ('Refunded', [BillingRecord.PaymentStatus.REFUNDED]),
)


def _billing_row(r):
    paid = r.status == BillingRecord.PaymentStatus.PAID
    return [
        r.transaction.order_number if r.transaction else '—',
        timezone.localtime(r.date_recorded),
        r.payment_method,
        r.reference_number or '—',
        r.status,
        (person_display_name(r.confirmed_by) or '—') if r.confirmed_by else '—',
        r.amount_due,
        r.amount_due if paid else None,
    ]


def build_billing_document(user, start, end, active_range='30', branch_id=None, now=None):
    records = list(scoped_billing_in_range(user, start, end, branch_id=branch_id))
    figures = billing_figures(records)

    paid = BillingRecord.PaymentStatus.PAID

    def branch_of(r):
        return r.transaction.branch if r.transaction else None

    # --- Summary, all from `records` ------------------------------------------
    status_rows = []
    for label, statuses in _BILLING_STATUS_ORDER:
        chosen = [r for r in records if r.status in statuses]
        status_rows.append([label, len(chosen), sum((r.amount_due for r in chosen), ZERO)])
    status_rows.append(['Total', len(records), sum((r.amount_due for r in records), ZERO)])

    methods = {name: [0, ZERO, ZERO] for name in PAYMENT_GROUP_ORDER}
    for r in records:
        bucket = methods[payment_group(r.payment_method)]
        bucket[0] += 1
        bucket[1] += r.amount_due
        if r.status == paid:
            bucket[2] += r.amount_due
    method_rows = [[name, *methods[name]] for name in PAYMENT_GROUP_ORDER]
    method_rows.append(['Total', len(records),
                        sum((m[1] for m in methods.values()), ZERO),
                        sum((m[2] for m in methods.values()), ZERO)])

    collected = _money(figures['collected'])
    headline = [
        ('Billing records', len(records)),
        ('Verified payments', figures['collected_count']),
        ('Pending payments', figures['pending_count']),
        ('Total amount verified', collected),
    ]
    summary = [
        SummaryBlock('Payment status', ['Status', 'Records', 'Amount due'], status_rows, total_row=True),
        SummaryBlock('By payment method', ['Method', 'Records', 'Amount due', 'Verified'],
                     method_rows, total_row=True),
        SummaryBlock('Revenue (completed orders)', ['Measure', 'Amount'], [
            ['Gross revenue', _money(figures['gross_revenue'])],
            ['Refunds (approved)', _money(figures['refund_value'])],
            ['Net revenue', _money(figures['total_revenue'])],
            ['Pending collection', _money(figures['outstanding'])],
        ]),
    ]

    # --- Detail ---------------------------------------------------------------
    buckets = _by_branch(records, branch_of)
    multi = len(buckets) > 1
    groups = []
    for branch, rows in buckets:
        group = Group(title=_group_title(branch), rows=[_billing_row(r) for r in rows])
        if multi:
            group.subtotal_label = f'Subtotal · {group.title} · {_plural(len(rows), "record")}'
            group.subtotal = [
                sum((r.amount_due for r in rows), ZERO),
                sum((r.amount_due for r in rows if r.status == paid), ZERO),
            ]
        groups.append(group)

    return ReportDocument(
        kind='billing', title='Billing Report', orientation='landscape',
        period_label=period_label(active_range, start, end),
        columns=BILLING_COLUMNS, label_span=6,
        headline=headline, summary=summary, groups=groups,
        grand_total=[sum((r.amount_due for r in records), ZERO), collected],
        grand_total_label=f'Grand total · {_plural(len(records), "record")}',
        notes=[
            'Verified = payment marked Paid. Refunded and partially refunded records are not counted '
            'as verified, matching the Billing Report page.',
            'Revenue is completed orders only, net of approved refunds — the same Net Sales as the Dashboard.',
            PAYMENT_GROUP_NOTE,
        ],
        detail_title='Billing records', sheet_name='Billing Detail',
        **_document_header(user, branch_id, now),
    )


# =============================================================================
# Inventory
# =============================================================================

INVENTORY_COLUMNS = [
    Column('sku', 'SKU', width=12),
    Column('product', 'Product', width=30),
    Column('on_hand', 'On hand', 'right', 8),
    Column('reorder', 'Reorder level', 'right', 9),
    Column('unit_cost', 'Unit cost (PHP)', 'right', 11),
    Column('value', 'Stock value (PHP)', 'right', 12),
    Column('status', 'Status', width=9),
    Column('expiry', 'Nearest expiry', width=10),
]

# Order a reader scans the status counts in: healthy first, worst last.
_STOCK_STATUS_ORDER = ('In Stock', 'Low Stock', 'Critical', 'No Stock')


def _inventory_branch_ids(user, branch_id):
    """The branch ids the inventory row set covers (None = every branch)."""
    chosen = owner_branch_choice(user, branch_id)
    if chosen:
        return [chosen]
    return branch_ids_for_user(user)


def build_inventory_document(user, branch_id=None, now=None):
    inventory_qs = scoped_inventory(user, branch_id=branch_id)
    records, figures = inventory_figures(inventory_qs)
    nearest = nearest_floor_expiry(_inventory_branch_ids(user, branch_id))
    today = timezone.localdate()

    # --- Summary, all from `records` ------------------------------------------
    status_of = {}
    near_expiry = {'Expiring Soon': 0, 'Approaching Expiry': 0}
    expired = 0
    for inv in records:
        status_of[inv.pk] = stock_status(inv.quantity_on_hand, inv.reorder_level)
        expiry = nearest.get((inv.product_id, inv.branch_id))
        # Same rule as the Non-Performing & Near-Expiry table (only rows that
        # still hold stock carry a floor batch, so quantity needs no check here).
        if expiry is not None and inv.quantity_on_hand > 0:
            tier = _expiry_tier((expiry - today).days)
            if tier in near_expiry:
                near_expiry[tier] += 1
            elif tier == 'Expired':
                expired += 1

    status_rows = []
    for name in _STOCK_STATUS_ORDER:
        chosen = [inv for inv in records if status_of[inv.pk] == name]
        status_rows.append([name, len(chosen), sum((_money(inv.value) for inv in chosen), ZERO)])
    status_rows.append(['Total', len(records), sum((row[2] for row in status_rows), ZERO)])
    near_total = sum(near_expiry.values())
    reorder_alerts = (
        sum(1 for s in status_of.values() if s in ('Critical', 'No Stock'))
    )

    total_value = _money(figures['total_value'])
    headline = [
        ('Total products', figures['product_lines']),
        ('Total stock value', total_value),
        ('Total units on hand', figures['total_units']),
        ('Near-expiry items', near_total),
    ]
    summary = [
        SummaryBlock('Stock status', ['Status', 'Stock lines', 'Stock value'], status_rows, total_row=True),
        SummaryBlock('Alerts (stock lines)', ['Measure', 'Lines'], [
            ['At or below reorder level (Critical + No Stock)', reorder_alerts],
            [f'Near-expiry, within {APPROACHING_EXPIRY_DAYS} days', near_total],
            [f'Expiring soon, within {EXPIRING_SOON_DAYS} days', near_expiry['Expiring Soon']],
            [f'Approaching expiry, {EXPIRING_SOON_DAYS + 1}-{APPROACHING_EXPIRY_DAYS} days',
             near_expiry['Approaching Expiry']],
            ['Already expired, still in stock', expired],
        ]),
    ]

    # --- Detail ---------------------------------------------------------------
    buckets = _by_branch(records, lambda inv: inv.branch)
    multi = len(buckets) > 1
    groups = []
    for branch, rows in buckets:
        group = Group(title=_group_title(branch), rows=[
            [
                inv.product.sku_code if inv.product else '—',
                inv.product.product_name if inv.product else '(product removed)',
                inv.quantity_on_hand,
                inv.reorder_level,
                _money(_inventory_unit_cost(inv)),
                _money(inv.value),
                status_of[inv.pk],
                nearest.get((inv.product_id, inv.branch_id)) if inv.quantity_on_hand > 0 else None,
            ]
            for inv in rows
        ])
        if multi:
            group.subtotal_label = f'Subtotal · {group.title} · {_plural(len(rows), "stock line")}'
            group.subtotal = [
                sum(inv.quantity_on_hand for inv in rows), None, None,
                sum((_money(inv.value) for inv in rows), ZERO), None, None,
            ]
        groups.append(group)

    return ReportDocument(
        kind='inventory', title='Inventory Report', orientation='landscape',
        period_label=f'Stock on hand as of {timezone.localtime(now or timezone.now()):%d %b %Y}',
        columns=INVENTORY_COLUMNS, label_span=2,
        headline=headline, summary=summary, groups=groups,
        grand_total=[figures['total_units'], None, None, total_value, None, None],
        grand_total_label=f'Grand total · {_plural(len(records), "stock line")}',
        notes=[
            'Stock value = units on hand x unit cost, where unit cost is the deepest bulk price tier '
            '(else the retail price) — the same basis as the Inventory Report page.',
            'Status: No Stock = 0 units; Critical = at or below the reorder level; Low Stock = up to 10% '
            'above it; In Stock = above that.',
            'Nearest expiry is the earliest expiry among the floor batches still holding stock.',
        ],
        detail_title='Stock detail', sheet_name='Inventory Detail',
        **_document_header(user, branch_id, now),
    )


# =============================================================================
# Procurement
# =============================================================================

PROCUREMENT_COLUMNS = [
    Column('po', 'PO No.', width=11),
    Column('supplier', 'Supplier', width=21),
    Column('ordered_on', 'Order date', width=10),
    Column('delivery', 'Delivery date', width=10),
    Column('status', 'Status', width=12),
    Column('flag', 'Receipt flag', width=9),
    Column('returned', 'Units returned', 'right', 8),
    Column('ordered', 'Ordered (PHP)', 'right', 10),
    Column('received', 'Received (PHP)', 'right', 10),
]

# Lifecycle order, so the status table reads the way a purchase order moves.
_PO_STATUS_ORDER = [choice.value for choice in Procurement.Status]


def _supplier_name(p):
    if not p.supplier:
        return '—'
    profile = getattr(p.supplier, 'supplier_profile', None)
    return getattr(profile, 'supplier_name', None) or p.supplier.username


def build_procurement_document(user, start, end, active_range='30', branch_id=None, now=None):
    procurement_qs = scoped_procurement_in_range(user, start, end, branch_id=branch_id)
    records = list(procurement_qs.select_related('supplier__supplier_profile').prefetch_related('items'))
    figures = procurement_figures(procurement_qs, records)
    cancelled = Procurement.Status.CANCELLED

    # --- Summary, all from the same queryset ----------------------------------
    returned_by_po = dict(
        SupplierReturnItem.objects.filter(
            supplier_return__procurement__in=procurement_qs,
            supplier_return__status__in=SupplierReturn.ACTIVE_STATUSES,
        ).values_list('supplier_return__procurement_id').annotate(units=Sum('quantity_returned'))
    )
    returns_raised = SupplierReturn.objects.filter(
        procurement__in=procurement_qs, status__in=SupplierReturn.ACTIVE_STATUSES,
    ).count()
    rejected_deliveries = DeliveryReceipt.objects.filter(
        procurement__in=procurement_qs, outcome=DeliveryReceipt.Outcome.REJECTED,
    ).count()
    discrepancy_pos = sum(1 for p in records if p.receipt_flag == Procurement.ReceiptFlag.DISCREPANCY)

    received_total = sum((_money(p.actual_received_cost) for p in records), ZERO)
    status_value = {}
    for p in records:
        status_value[p.status] = status_value.get(p.status, ZERO) + p.total_cost
    status_rows = [
        [name, figures['status_counts'][name], status_value.get(name, ZERO)]
        for name in _PO_STATUS_ORDER if figures['status_counts'].get(name)
    ]
    status_rows.append(['Total', len(records), sum((row[2] for row in status_rows), ZERO)])

    returns = figures['return_totals']
    gross_procurement = _money(figures['gross_procurement'])
    headline = [
        ('Purchase orders', len(records)),
        ('Total ordered value', gross_procurement),
        ('Total received value', received_total),
        ('Net procurement', _money(figures['net_procurement'])),
    ]
    summary = [
        SummaryBlock('Purchase orders by status', ['Status', 'POs', 'Ordered value'],
                     status_rows, total_row=True),
        SummaryBlock('Value', ['Measure', 'Amount'], [
            ['Total ordered (excl. cancelled)', gross_procurement],
            ['Total received (at cost)', received_total],
            ['Closed POs (fully received)', _money(figures['received_value'])],
            ['Still open (not yet received)', _money(figures['active_value'])],
            ['Net procurement (after returns)', _money(figures['net_procurement'])],
        ]),
        SummaryBlock('Rejected and returned', ['Measure', 'Count / amount'], [
            ['Deliveries rejected on receipt', rejected_deliveries],
            ['POs with a receipt discrepancy', discrepancy_pos],
            ['Supplier returns raised', returns_raised],
            ['Units returned to suppliers', sum(returned_by_po.values())],
            ['Value of goods shipped back', _money(returns['returned_goods'])],
            ['Supplier credit notes settled', _money(returns['credits'])],
        ]),
    ]

    # --- Detail ---------------------------------------------------------------
    def row(p):
        flag = None if p.receipt_flag == Procurement.ReceiptFlag.NONE else p.receipt_flag
        return [
            p.po_number, _supplier_name(p), p.order_date, p.delivery_date, p.status,
            flag, returned_by_po.get(p.pk) or None, _money(p.total_cost), _money(p.actual_received_cost),
        ]

    buckets = _by_branch(records, lambda p: p.branch)
    multi = len(buckets) > 1
    groups = []
    for branch, rows in buckets:
        group = Group(title=_group_title(branch), rows=[row(p) for p in rows])
        if multi:
            group.subtotal_label = f'Subtotal · {group.title} · {_plural(len(rows), "purchase order")}'
            scoped = procurement_qs.filter(branch=branch) if branch else procurement_qs.filter(branch__isnull=True)
            group.subtotal = [
                sum(returned_by_po.get(p.pk, 0) for p in rows),
                procurement_cost_total(scoped),
                sum((_money(p.actual_received_cost) for p in rows), ZERO),
            ]
        groups.append(group)

    return ReportDocument(
        kind='procurement', title='Procurement Report', orientation='landscape',
        period_label=period_label(active_range, start, end),
        columns=PROCUREMENT_COLUMNS, label_span=6,
        headline=headline, summary=summary, groups=groups,
        grand_total=[sum(returned_by_po.values()), gross_procurement, received_total],
        grand_total_label=f'Grand total · {_plural(len(records), "purchase order")}',
        notes=[
            f'Ordered value counts every purchase order except {cancelled}. The amount shown on a '
            f'{cancelled} order is for reference only. Received value is what physically arrived, at the '
            'confirmed unit cost, and includes partial deliveries.',
            'Units returned counts supplier returns that are requested, acknowledged, shipped or settled.',
        ],
        detail_title='Purchase orders', sheet_name='Procurement Detail',
        **_document_header(user, branch_id, now),
    )


BUILDERS = {
    'sales': build_sales_document,
    'billing': build_billing_document,
    'inventory': build_inventory_document,
    'procurement': build_procurement_document,
}


def build_document(user, kind, start=None, end=None, active_range='30', branch_id=None, now=None):
    """The ReportDocument for `kind`, scoped to what `user` may see. Inventory
    is a live snapshot and ignores the date range."""
    if kind == 'inventory':
        return build_inventory_document(user, branch_id=branch_id, now=now)
    return BUILDERS[kind](user, start, end, active_range=active_range, branch_id=branch_id, now=now)


# =============================================================================
# PDF
# =============================================================================

class ExportTooLarge(Exception):
    """The report has more rows than a PDF can hold in reasonable time."""


def _pdf_context(doc):
    total = sum(c.width for c in doc.columns)
    columns = [
        {'label': c.label, 'align': c.align, 'width': round(c.width * 100 / total, 1)}
        for c in doc.columns
    ]

    def cells(values):
        return [
            {'text': fmt_cell(v), 'align': col.align}
            for v, col in zip(values, doc.columns[len(doc.columns) - len(values):])
        ]

    span = doc.label_span
    groups = []
    for g in doc.groups:
        groups.append({
            'title': g.title,
            'rows': [cells(r) for r in g.rows],
            'subtotal': (
                {'label': g.subtotal_label, 'span': span, 'cells': cells(g.subtotal)}
                if g.subtotal is not None else None
            ),
        })
    grand_total = None
    if doc.grand_total is not None:
        grand_total = {'label': doc.grand_total_label, 'span': span, 'cells': cells(doc.grand_total)}

    def block(b):
        # Left-align the first column, right-align the figures. The label
        # column takes the room a short block leaves it, and an amount gets
        # more than a count, so labels and "PHP 12,345.00" both stay on a line.
        weights, is_money = [5], [False]
        for i in range(1, len(b.headers)):
            sample = next((row[i] for row in b.rows if row[i] is not None), None)
            is_money.append(isinstance(sample, Decimal))
            weights.append(3 if is_money[-1] else 2)
        scale = 100 / sum(weights)
        # Two columns leave room for "PHP 12,345.00"; with three or four the
        # amounts would wrap, so the heading carries the currency instead.
        spell_out_php = len(b.headers) == 2
        return {
            'title': b.title,
            'headers': [
                {
                    'label': label if spell_out_php or not is_money[i] else f'{label} (PHP)',
                    'width': round(weights[i] * scale),
                    'align': 'left' if i == 0 else 'right',
                }
                for i, label in enumerate(b.headers)
            ],
            'rows': [
                [{'text': fmt_cell(v, money_prefix=spell_out_php), 'align': 'left' if i == 0 else 'right'}
                 for i, v in enumerate(row)]
                for row in b.rows
            ],
            'total_row': b.total_row,
        }

    blocks = [block(b) for b in doc.summary]
    # Up to three summary tables side by side; the cell width follows the count.
    block_rows = []
    for i in range(0, len(blocks), 3):
        chunk = blocks[i:i + 3]
        block_rows.append([{**b, 'width': round(100 / len(chunk))} for b in chunk])

    # A text tile ("Plaridel Branch · PHP 14,088.00") needs more room than a figure.
    headline_weights = [1.7 if isinstance(value, str) else 1 for _, value in doc.headline]

    # A4 is 297 x 210 mm and the page margins are 11 mm a side. The footer is
    # two static frames side by side (running text left, "Page X of Y" right),
    # so the left one takes whatever the 52 mm right one leaves.
    usable_mm = (297 if doc.orientation == 'landscape' else 210) - 22

    return {
        'doc': doc,
        'store_name': STORE_NAME,
        'app_name': APP_NAME,
        'page_size': f'A4 {doc.orientation}',
        'footer_left_mm': usable_mm - 52 - 2,
        'columns': columns,
        'groups': groups,
        'grand_total': grand_total,
        'headline': [
            {'label': label, 'text': fmt_cell(value, money_prefix=True),
             'width': round(weight * 100 / sum(w for w in headline_weights))}
            for (label, value), weight in zip(doc.headline, headline_weights)
        ],
        'block_rows': block_rows,
    }


def render_pdf(doc):
    """The report as PDF bytes. Raises ExportTooLarge past MAX_PDF_ROWS, and
    RuntimeError if xhtml2pdf cannot lay the document out."""
    from xhtml2pdf import pisa

    if doc.row_count > MAX_PDF_ROWS:
        raise ExportTooLarge(
            f'This report has {doc.row_count:,} rows, more than a PDF can hold ({MAX_PDF_ROWS:,}). '
            f'Narrow the date range or pick one branch, or use Export Excel.'
        )
    html = render_to_string('reports/pdf/report.html', _pdf_context(doc))
    buffer = io.BytesIO()
    result = pisa.CreatePDF(src=html, dest=buffer)
    if result.err:
        raise RuntimeError('The PDF report could not be generated.')
    return buffer.getvalue()


# =============================================================================
# Excel
# =============================================================================

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
_XLSX_MONEY = '#,##0.00'
_XLSX_MONEY_SUMMARY = '"PHP" #,##0.00'
_NO_RECORDS = 'No records for the selected filters'


def _xlsx_value(cell, value, money_format=_XLSX_MONEY):
    """Write `value` to an openpyxl cell as the right Excel type.

    openpyxl turns any string starting with "=" into a formula, which would
    execute for whoever opens the file (a customer-supplied name of
    =cmd|'/c calc'!A1) — such strings are stored as literal text instead."""
    if value is None or value == '':
        cell.value = None
        return
    if isinstance(value, Decimal):
        cell.value = float(value)
        cell.number_format = money_format
    elif isinstance(value, datetime):
        cell.value = timezone.make_naive(value) if timezone.is_aware(value) else value
        cell.number_format = 'yyyy-mm-dd hh:mm'
    elif isinstance(value, date):
        cell.value = value
        cell.number_format = 'yyyy-mm-dd'
    elif isinstance(value, int):
        cell.value = value
        cell.number_format = '#,##0'
    else:
        cell.value = str(value)
        if cell.value.startswith('='):
            cell.data_type = 's'


def render_xlsx(doc):
    """The report as .xlsx bytes: a "Summary" sheet with the same header and
    Summary as the PDF, then the detail sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    green = '1E7A4F'
    head_fill = PatternFill('solid', start_color=green)
    soft_fill = PatternFill('solid', start_color='E3F3EB')
    grey_fill = PatternFill('solid', start_color='F2F2F2')
    white_bold = Font(bold=True, color='FFFFFF')
    bold = Font(bold=True)
    top_rule = Border(top=Side(style='thin', color='333333'))

    wb = Workbook()

    # --- Summary sheet --------------------------------------------------------
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = STORE_NAME
    ws['A1'].font = Font(bold=True, size=15)
    ws['A2'] = APP_NAME
    ws['A2'].font = Font(color='555555')
    ws['A3'] = doc.title
    ws['A3'].font = Font(bold=True, size=13, color=green)
    meta = [
        ('Branch', doc.branch_label), ('Period', doc.period_label),
        ('Generated by', doc.generated_by), ('Generated', doc.generated_label),
    ]
    row = 5
    for label, value in meta:
        ws.cell(row=row, column=1, value=label).font = bold
        _xlsx_value(ws.cell(row=row, column=2), value)
        row += 1
    row += 1

    if doc.is_empty:
        ws.cell(row=row, column=1, value=_NO_RECORDS).font = Font(bold=True, size=12, color='B42318')
    else:
        ws.cell(row=row, column=1, value='Summary').font = Font(bold=True, size=13)
        row += 1
        for col, text in enumerate(('Key figures', ''), start=1):
            ws.cell(row=row, column=col, value=text or None).font = white_bold
            ws.cell(row=row, column=col).fill = head_fill
        row += 1
        for label, value in doc.headline:
            ws.cell(row=row, column=1, value=label)
            _xlsx_value(ws.cell(row=row, column=2), value, _XLSX_MONEY_SUMMARY)
            ws.cell(row=row, column=2).alignment = Alignment(horizontal='right')
            row += 1
        row += 1

        for block in doc.summary:
            ws.cell(row=row, column=1, value=block.title).font = Font(bold=True, color=green)
            row += 1
            for col, text in enumerate(block.headers, start=1):
                c = ws.cell(row=row, column=col, value=text)
                c.font, c.fill = white_bold, head_fill
                if col > 1:
                    c.alignment = Alignment(horizontal='right')
            row += 1
            for i, values in enumerate(block.rows):
                last = block.total_row and i == len(block.rows) - 1
                for col, value in enumerate(values, start=1):
                    c = ws.cell(row=row, column=col)
                    _xlsx_value(c, value.strip() if isinstance(value, str) else value, _XLSX_MONEY_SUMMARY)
                    if col > 1:
                        c.alignment = Alignment(horizontal='right')
                    if last:
                        c.font, c.fill, c.border = bold, grey_fill, top_rule
                row += 1
            row += 1

        for note in doc.notes:
            c = ws.cell(row=row, column=1, value=note)
            c.font = Font(italic=True, color='555555', size=9)
            row += 1

    ws.column_dimensions['A'].width = 52
    for letter in 'BCD':
        ws.column_dimensions[letter].width = 22
    ws.sheet_view.showGridLines = False

    # --- Detail sheet ---------------------------------------------------------
    ds = wb.create_sheet(doc.sheet_name)
    ds['A1'] = f'{doc.title} — {doc.detail_title}'
    ds['A1'].font = Font(bold=True, size=13, color=green)
    ds['A2'] = (f'{doc.branch_label} · {doc.period_label} · Generated by {doc.generated_by} '
                f'on {doc.generated_label}')
    ds['A2'].font = Font(color='555555', size=9)

    header_row = 4
    for col, column in enumerate(doc.columns, start=1):
        c = ds.cell(row=header_row, column=col, value=column.label)
        c.font, c.fill = white_bold, head_fill
        c.alignment = Alignment(horizontal=column.align, vertical='center', wrap_text=True)
    ds.row_dimensions[header_row].height = 30

    row = header_row + 1
    if doc.is_empty:
        ds.cell(row=row, column=1, value=_NO_RECORDS).font = Font(bold=True, color='B42318')
    else:
        width = len(doc.columns)

        def write_total(label, values, fill):
            nonlocal row
            ds.cell(row=row, column=1, value=label)
            if doc.label_span > 1:
                ds.merge_cells(start_row=row, start_column=1, end_row=row, end_column=doc.label_span)
            for offset, value in enumerate(values):
                _xlsx_value(ds.cell(row=row, column=doc.label_span + 1 + offset), value)
            for col in range(1, width + 1):
                c = ds.cell(row=row, column=col)
                c.font, c.fill, c.border = bold, fill, top_rule
                if col > doc.label_span:
                    c.alignment = Alignment(horizontal='right')
            row += 1

        for group in doc.groups:
            title = ds.cell(row=row, column=1, value=f'Branch: {group.title}')
            ds.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
            title.font, title.fill = Font(bold=True, color='14532D'), soft_fill
            row += 1
            for values in group.rows:
                for col, (value, column) in enumerate(zip(values, doc.columns), start=1):
                    c = ds.cell(row=row, column=col)
                    _xlsx_value(c, value)
                    if column.align != 'left':
                        c.alignment = Alignment(horizontal=column.align)
                row += 1
            if group.subtotal is not None:
                write_total(group.subtotal_label, group.subtotal, grey_fill)
            row += 1
        if doc.grand_total is not None:
            write_total(doc.grand_total_label, doc.grand_total, soft_fill)

        row += 1
        for note in doc.notes:
            ds.cell(row=row, column=1, value=note).font = Font(italic=True, color='555555', size=9)
            row += 1

    total_weight = sum(c.width for c in doc.columns)
    for col, column in enumerate(doc.columns, start=1):
        ds.column_dimensions[get_column_letter(col)].width = max(12, round(column.width * 140 / total_weight))
    ds.freeze_panes = ds.cell(row=header_row + 1, column=1)
    ds.print_title_rows = f'{header_row}:{header_row}'
    ds.page_setup.orientation = doc.orientation
    ds.page_setup.fitToWidth = 1
    ds.page_setup.fitToHeight = 0
    ds.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = 'portrait'

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
