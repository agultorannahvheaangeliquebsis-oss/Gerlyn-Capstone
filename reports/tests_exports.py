"""Tests for the PDF / Excel report exports (reports/exports.py).

The point of these exports is that the Summary opening each file can be
trusted. So the central tests build one multi-branch fixture and assert that
every Summary figure

  * is the same number the report page and the dashboard show for the same
    filters (the page/dashboard context is read through the Django test client,
    not recomputed here), and
  * is the total of the detail rows the same file lists — rows, per-branch
    subtotals and the grand total all reconcile with the Summary.

The PDF tests read the real bytes back with pypdf (installed as a dependency of
xhtml2pdf, which the purchase-order and receipt PDFs already use): page count,
"Page X of Y", repeated table headings, and that every row lands on exactly one
page — the "no cut-off rows" requirement. A test fails if xhtml2pdf logs that a
static frame dropped content, because that is how an overflowing header shows
up.
"""
import io
import re
from datetime import timedelta
from decimal import Decimal
from unittest import mock

from django.test import Client, SimpleTestCase, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader

from accounts.models import User
from billing.models import BillingRecord
from branches.models import Branch
from inventory.alerts import stock_status
from inventory.models import Batch, Inventory, Product, ProductCategory
from procurement.models import (
    DeliveryReceipt, Procurement, ProcurementItem, SupplierReturn, SupplierReturnItem,
)
from reports import exports
from reports.views import _non_performing_products
from sales.models import Refund, SalesTransaction

D = Decimal


def pdf_pages(response):
    """The text of each page of a PDF response."""
    assert response['Content-Type'] == 'application/pdf'
    reader = PdfReader(io.BytesIO(response.content))
    return [page.extract_text() or '' for page in reader.pages]


def pdf_reader(response):
    return PdfReader(io.BytesIO(response.content))


def block_rows(doc, title):
    """{first-column label: remaining cells} of one Summary block."""
    block = next(b for b in doc.summary if b.title.startswith(title))
    return {row[0]: row[1:] for row in block.rows}


def headline(doc):
    return dict(doc.headline)


class ExportFixture(TestCase):
    """One owner, two branches and a manager, with sales / billing / stock /
    purchase orders chosen so every Summary line has a distinct value."""

    @classmethod
    def setUpTestData(cls):
        cls.today = timezone.localdate()
        cls.start = cls.today - timedelta(days=6)
        cls.owner = User.objects.create_user(
            username='exp-owner', password='pw', role=User.Role.OWNER,
            first_name='Olivia', last_name='Owner',
        )
        cls.branch_a = Branch.objects.create(branch_name='Export Branch A', location='Baliuag', owner=cls.owner)
        cls.branch_b = Branch.objects.create(branch_name='Export Branch B', location='Pulilan', owner=cls.owner)

        cls.manager = User.objects.create_user(
            username='exp-manager', password='pw', role=User.Role.BRANCH_MANAGER,
        )
        cls.manager.manager_profile.branch = cls.branch_a
        cls.manager.manager_profile.save(update_fields=['branch'])

        cls.customer = User.objects.create_user(
            username='exp-customer', password='pw', role=User.Role.CUSTOMER,
            first_name='Cora', last_name='Customer',
        )

        cls._build_sales()
        cls._build_billing()
        cls._build_inventory()
        cls._build_procurement()

    # --- sales ----------------------------------------------------------------

    @classmethod
    def _sale(cls, branch, amount, status, mode, channel='Physical', customer=None, when=None):
        txn = SalesTransaction.objects.create(
            branch=branch, total_amount=D(amount), transaction_type=channel,
            mode_of_payment=mode, customer=customer, recorded_by=cls.owner,
        )
        updates = {}
        if txn.order_status != status:
            updates['order_status'] = status
        if when is not None:
            updates['transaction_date'] = when
        if updates:
            SalesTransaction.objects.filter(pk=txn.pk).update(**updates)
            txn.refresh_from_db()
        return txn

    @classmethod
    def _build_sales(cls):
        S = SalesTransaction.OrderStatus
        a, b = cls.branch_a, cls.branch_b
        cls.a1 = cls._sale(a, '1000.00', S.COMPLETED, 'Cash')
        cls.a2 = cls._sale(a, '500.00', S.COMPLETED, 'GCash', 'Online', cls.customer)
        cls.a3 = cls._sale(a, '300.00', S.COMPLETED, 'Card')
        cls.a4 = cls._sale(a, '200.00', S.COMPLETED, 'Maya', 'Online', cls.customer)
        cls.a5 = cls._sale(a, '150.00', S.COMPLETED, 'Cash on Delivery')
        cls.a6 = cls._sale(a, '700.00', S.CANCELLED, 'Cash')
        cls.a7 = cls._sale(a, '400.00', S.PENDING, 'GCash', 'Online', cls.customer)
        cls.b1 = cls._sale(b, '250.00', S.COMPLETED, 'Cash')
        cls.b2 = cls._sale(b, '350.00', S.COMPLETED, 'QR Ph', 'Online', cls.customer)
        cls.b3 = cls._sale(b, '90.00', S.PREPARING, 'Cash', 'Online', cls.customer)
        # Outside the 7-day range: must not reach any figure.
        cls.old = cls._sale(a, '9999.00', S.COMPLETED, 'Cash', when=timezone.now() - timedelta(days=40))

        def refund(status, amount):
            Refund.objects.create(
                transaction=cls.a1, reason='t', reason_code=Refund.ReasonCode.OTHER,
                refund_type=Refund.RefundType.PARTIAL, status=status,
                refund_method=Refund.RefundMethod.CASH, total_amount=D(amount),
                requested_by=cls.owner, refunded_by=cls.owner, refunded_at=timezone.now(),
            )
        refund(Refund.Status.APPROVED, '100.00')
        # Requests that moved no money: neither may reach a refund figure.
        refund(Refund.Status.PENDING, '75.00')
        refund(Refund.Status.REJECTED, '60.00')

        cls.in_range = [cls.a1, cls.a2, cls.a3, cls.a4, cls.a5, cls.a6, cls.a7, cls.b1, cls.b2, cls.b3]
        cls.expect = {
            'orders': 10, 'completed': 7, 'cancelled': 1, 'in_progress': 2,
            'walk_in': 5, 'online': 5, 'refunded_orders': 1,
            'gross': D('2750.00'), 'refunds': D('100.00'), 'net': D('2650.00'),
            'pipeline': D('490.00'),
        }

    # --- billing --------------------------------------------------------------

    @classmethod
    def _build_billing(cls):
        P = BillingRecord.PaymentStatus
        # One record per sale (the post_save signal). Verified (Paid): a1 a2 a3 b1.
        # Pending: the rest. Methods follow each sale's own mode.
        for txn in cls.in_range + [cls.old]:
            record = txn.billing_record
            fields = {'payment_method': txn.mode_of_payment}
            if txn in (cls.a1, cls.a2, cls.a3, cls.b1):
                fields.update(status=P.PAID, confirmed_by=cls.owner)
            if txn == cls.old:
                fields['date_recorded'] = timezone.now() - timedelta(days=40)
            BillingRecord.objects.filter(pk=record.pk).update(**fields)
        cls.verified_amount = D('1000.00') + D('500.00') + D('300.00') + D('250.00')

    # --- inventory ------------------------------------------------------------

    @classmethod
    def _build_inventory(cls):
        category = ProductCategory.objects.create(category_name='Export Category')
        cls.products = []
        # (qty, reorder) -> No Stock / Critical / Low Stock / In Stock, per branch.
        plan = [(0, 10), (5, 10), (10, 10), (11, 10), (100, 10)]
        for i, (qty, reorder) in enumerate(plan):
            product = Product.objects.create(
                category=category, sku_code=f'SKU-EXP-{i}', product_name=f'Export Product {i}',
                unit_price=D(10 + i),
            )
            cls.products.append(product)
            Inventory.objects.create(branch=cls.branch_a, product=product, quantity_on_hand=qty, reorder_level=reorder)
        # Branch B carries two of the same products (the product count is distinct, not per branch).
        Inventory.objects.create(branch=cls.branch_b, product=cls.products[0], quantity_on_hand=40, reorder_level=10)
        Inventory.objects.create(branch=cls.branch_b, product=cls.products[1], quantity_on_hand=0, reorder_level=10)

        def batch(product, branch, qty, days):
            Batch.objects.create(
                branch=branch, product=product, quantity=qty, location=Batch.Location.FLOOR,
                received_date=cls.today, expiry_date=cls.today + timedelta(days=days),
            )
        batch(cls.products[1], cls.branch_a, 5, 10)     # expiring soon
        batch(cls.products[3], cls.branch_a, 11, 60)    # approaching
        batch(cls.products[4], cls.branch_a, 100, 400)  # fresh
        batch(cls.products[0], cls.branch_b, 40, -3)    # expired, still on the shelf

    # --- procurement ----------------------------------------------------------

    @classmethod
    def _build_procurement(cls):
        cls.supplier = User.objects.create_user(username='exp-supplier', password='pw', role=User.Role.SUPPLIER)
        Status = Procurement.Status

        def po(branch, status, cost, days_ago=1, flag='None'):
            return Procurement.objects.create(
                branch=branch, supplier=cls.supplier, status=status, total_cost=D(cost),
                order_date=cls.today - timedelta(days=days_ago), receipt_flag=flag,
            )
        cls.po_received = po(cls.branch_a, Status.RECEIVED, '1000.00')
        cls.po_partial = po(cls.branch_a, Status.PARTIALLY_RECEIVED, '600.00', flag='Discrepancy')
        cls.po_approved = po(cls.branch_a, Status.APPROVED, '300.00')
        cls.po_cancelled = po(cls.branch_b, Status.CANCELLED, '9999.00')
        cls.po_incoming = po(cls.branch_b, Status.INCOMING, '450.00')
        cls.po_old = po(cls.branch_a, Status.RECEIVED, '5000.00', days_ago=60)

        product = cls.products[0]
        item = ProcurementItem.objects.create(
            procurement=cls.po_received, product=product, quantity_ordered=20,
            quantity_received=20, unit_cost_snapshot=D('50.00'),
        )
        ProcurementItem.objects.create(
            procurement=cls.po_partial, product=product, quantity_ordered=12,
            quantity_received=5, unit_cost_snapshot=D('50.00'),
        )
        returned = SupplierReturn.objects.create(
            procurement=cls.po_received, supplier=cls.supplier, branch=cls.branch_a,
            reason=SupplierReturn.Reason.DAMAGED, status=SupplierReturn.Status.SHIPPED,
            requested_by=cls.owner,
        )
        SupplierReturnItem.objects.create(
            supplier_return=returned, procurement_item=item, product=product, quantity_returned=4,
        )
        DeliveryReceipt.objects.create(
            procurement=cls.po_incoming, outcome=DeliveryReceipt.Outcome.REJECTED,
            rejection_reason='Wrong goods',
        )

    # --- helpers --------------------------------------------------------------

    @property
    def range_params(self):
        return {'range': 'custom', 'date_from': self.start.isoformat(), 'date_to': self.today.isoformat()}

    def client_for(self, user):
        client = Client()
        client.force_login(user)
        return client

    def doc(self, kind, user=None, **extra):
        """The ReportDocument exactly as the export view builds it."""
        user = user or self.owner
        if kind == 'inventory':
            return exports.build_inventory_document(user, **extra)
        return exports.build_document(
            user, kind, self.start, self.today, active_range='custom', **extra,
        )


# =============================================================================
# The Summary must be the same number as the page, the dashboard and the rows
# =============================================================================

class PaymentGroupTests(SimpleTestCase):
    """Which Summary row each recorded payment method lands in."""

    def test_every_method_maps_to_one_of_the_summary_rows(self):
        expected = {
            'Cash': exports.CASH, 'Cash on Delivery': exports.CASH, 'Pay at Store': exports.CASH,
            'GCash': exports.GCASH, 'Maya': exports.MAYA, 'Card': exports.CARD,
            'QR Ph': exports.OTHER_METHODS, 'Charge Sale': exports.OTHER_METHODS,
            'Bank Transfer': exports.OTHER_METHODS,
        }
        # Every method either model can hold — a new one added later must be placed on purpose
        # (it would otherwise fall into Other methods, which this test then makes visible).
        every_method = set(SalesTransaction.ModeOfPayment.values) | set(BillingRecord.PaymentMethod.values)
        self.assertEqual(every_method, set(expected))
        for method, group in expected.items():
            with self.subTest(method=method):
                self.assertEqual(exports.payment_group(method), group)
                self.assertIn(group, exports.PAYMENT_GROUP_ORDER)


class SalesSummaryReconcilesTests(ExportFixture):

    def test_summary_matches_the_report_pages_and_the_dashboard(self):
        client = self.client_for(self.owner)
        doc = self.doc('sales')
        e = self.expect

        hub = client.get(reverse('reports_dashboard'), self.range_params).context       # Reports hub
        page = client.get(reverse('sales_report'), self.range_params).context['totals']  # Sales Report page
        dash = client.get(reverse('accounts:owner_dashboard'), self.range_params).context  # Dashboard
        cards = {c['key']: c for c in dash['kpi_cards']}

        head = headline(doc)
        # Against the fixture's own arithmetic...
        self.assertEqual(head['Total orders'], e['orders'])
        self.assertEqual(head['Gross sales'], e['gross'])
        self.assertEqual(head['Total refunds'], e['refunds'])
        self.assertEqual(head['Net sales'], e['net'])
        # ...against the Reports page...
        self.assertEqual(head['Gross sales'], hub['gross_sales'])
        self.assertEqual(head['Total refunds'], hub['refund_value'])
        self.assertEqual(head['Net sales'], hub['total_sales'])
        self.assertEqual(block_rows(doc, 'Orders')['Completed'][0], hub['total_orders'])
        self.assertEqual(block_rows(doc, 'Orders')['Online'][0], hub['online_orders'])
        self.assertEqual(block_rows(doc, 'Sales')['Committed, not yet sales (in progress)'][0], hub['pipeline_total'])
        self.assertEqual(head['Total orders'], page['transactions'])
        self.assertEqual(head['Net sales'], page['net_sales'])
        # ...and against the Dashboard. Its "Total Sales" card is gross of refunds
        # (the Reports headline is net); both figures are printed, labelled.
        self.assertEqual(cards['total_sales']['value_display'], f'₱{head["Gross sales"]:,.2f}')
        self.assertEqual(cards['total_transactions']['value_display'], f'{block_rows(doc, "Orders")["Completed"][0]:,}')

    def test_order_counts_partition_the_total(self):
        doc = self.doc('sales')
        orders = block_rows(doc, 'Orders')
        e = self.expect
        self.assertEqual(orders['Completed'][0], e['completed'])
        self.assertEqual(orders['Cancelled'][0], e['cancelled'])
        self.assertEqual(orders['In progress (pending, approved, preparing)'][0], e['in_progress'])
        self.assertEqual(orders['Refunded (with an approved refund)'][0], e['refunded_orders'])
        self.assertEqual(orders['Walk-in'][0], e['walk_in'])
        self.assertEqual(orders['Online'][0], e['online'])
        self.assertEqual(
            orders['Completed'][0] + orders['Cancelled'][0] + orders['In progress (pending, approved, preparing)'][0],
            orders['Total orders'][0],
        )
        self.assertEqual(orders['Walk-in'][0] + orders['Online'][0], orders['Total orders'][0])

    def test_payment_methods_group_and_add_up_to_gross_sales(self):
        doc = self.doc('sales')
        pay = block_rows(doc, 'Payments received')
        # Cash a1 1000 + b1 250, and Cash on Delivery a5 150 counts as cash.
        self.assertEqual(pay[exports.CASH], [3, D('1400.00')])
        # GCash and Maya each keep their own row. The pending GCash sale (a7) is not a payment.
        self.assertEqual(pay[exports.GCASH], [1, D('500.00')])
        self.assertEqual(pay[exports.MAYA], [1, D('200.00')])
        self.assertEqual(pay[exports.CARD], [1, D('300.00')])
        self.assertEqual(pay[exports.OTHER_METHODS], [1, D('350.00')])             # QR Ph b2
        self.assertEqual(sum(pay[g][1] for g in exports.PAYMENT_GROUP_ORDER), headline(doc)['Gross sales'])
        self.assertEqual(sum(pay[g][0] for g in exports.PAYMENT_GROUP_ORDER), self.expect['completed'])
        self.assertEqual(pay['Total'], [self.expect['completed'], headline(doc)['Gross sales']])

    def test_best_branch_only_when_every_branch_is_in_scope(self):
        doc = self.doc('sales')
        self.assertEqual(headline(doc)['Best branch'], 'Export Branch A · PHP 2,050.00')
        # One branch picked: nothing to rank.
        self.assertNotIn('Best branch', headline(self.doc('sales', branch_id=self.branch_b.pk)))
        # A Branch Manager only ever has their own branch.
        self.assertNotIn('Best branch', headline(self.doc('sales', user=self.manager)))

    def test_detail_rows_subtotals_and_grand_total_reconcile_with_the_summary(self):
        doc = self.doc('sales')
        head = headline(doc)
        rows = [r for g in doc.groups for r in g.rows]
        cols = [c.key for c in doc.columns]
        col = lambda r, key: r[cols.index(key)]  # noqa: E731

        # Exactly the orders in range, once each.
        self.assertEqual(sorted(col(r, 'order') for r in rows), sorted(t.order_number for t in self.in_range))
        self.assertEqual(len(rows), head['Total orders'])

        completed = [r for r in rows if col(r, 'status') == 'Completed']
        self.assertEqual(sum(col(r, 'amount') for r in completed), head['Gross sales'])
        self.assertEqual(sum((col(r, 'refunded') or D('0')) for r in completed), head['Total refunds'])
        self.assertEqual(sum(col(r, 'net') for r in completed), head['Net sales'])
        # Cancelled / in-progress rows carry no Net and are not totalled.
        self.assertTrue(all(col(r, 'net') is None for r in rows if r not in completed))

        self.assertEqual(doc.grand_total, [head['Gross sales'], head['Total refunds'], head['Net sales']])
        # Per-branch subtotals add up to the grand total.
        self.assertEqual([g.title for g in doc.groups], ['Export Branch A', 'Export Branch B'])
        for i in range(3):
            self.assertEqual(sum(g.subtotal[i] for g in doc.groups), doc.grand_total[i])
        a, b = doc.groups
        self.assertEqual(a.subtotal, [D('2150.00'), D('100.00'), D('2050.00')])
        self.assertEqual(b.subtotal, [D('600.00'), D('0.00'), D('600.00')])

    def test_orders_whose_branch_was_deleted_get_their_own_last_group(self):
        # Branch.on_delete=SET_NULL leaves such sales with no branch; they must still
        # be listed and totalled, not silently dropped from the subtotals.
        orphan = self._sale(None, '40.00', SalesTransaction.OrderStatus.COMPLETED, 'Cash')
        doc = self.doc('sales')
        self.assertEqual([g.title for g in doc.groups], ['Export Branch A', 'Export Branch B', 'No branch assigned'])
        self.assertEqual(doc.groups[-1].subtotal, [D('40.00'), D('0.00'), D('40.00')])
        self.assertEqual(headline(doc)['Gross sales'], self.expect['gross'] + D('40.00'))
        for i in range(3):
            self.assertEqual(sum(g.subtotal[i] for g in doc.groups), doc.grand_total[i])
        self.assertIn(orphan.order_number, [r[0] for g in doc.groups for r in g.rows])
        # ...and best branch ranks real branches only.
        self.assertEqual(headline(doc)['Best branch'], 'Export Branch A · PHP 2,050.00')

    def test_single_branch_report_has_a_grand_total_but_no_repeated_subtotal(self):
        doc = self.doc('sales', user=self.manager)
        self.assertEqual(len(doc.groups), 1)
        self.assertIsNone(doc.groups[0].subtotal)
        self.assertEqual(doc.grand_total[0], D('2150.00'))


class BillingSummaryReconcilesTests(ExportFixture):

    def test_summary_matches_the_billing_page_and_the_dashboard(self):
        client = self.client_for(self.owner)
        doc = self.doc('billing')
        page = client.get(reverse('billing_report'), self.range_params).context
        hub = client.get(reverse('reports_dashboard'), self.range_params).context
        head = headline(doc)

        self.assertEqual(head['Billing records'], 10)
        self.assertEqual(head['Verified payments'], 4)
        self.assertEqual(head['Pending payments'], 6)
        self.assertEqual(head['Total amount verified'], self.verified_amount)

        self.assertEqual(head['Verified payments'], page['collected_count'])
        self.assertEqual(head['Pending payments'], page['pending_count'])
        self.assertEqual(head['Total amount verified'], page['collected'])
        rev = block_rows(doc, 'Revenue')
        self.assertEqual(rev['Net revenue'][0], page['total_revenue'])
        self.assertEqual(rev['Pending collection'][0], page['outstanding'])
        # Net revenue is the Dashboard's net sales for the same filters.
        self.assertEqual(rev['Net revenue'][0], hub['total_sales'])
        self.assertEqual(rev['Gross revenue'][0], hub['gross_sales'])
        self.assertEqual(rev['Refunds (approved)'][0], hub['refund_value'])

    def test_status_and_method_tables_account_for_every_record(self):
        doc = self.doc('billing')
        status = block_rows(doc, 'Payment status')
        self.assertEqual(status['Verified (Paid)'], [4, self.verified_amount])
        self.assertEqual(status['Pending'][0], 6)
        self.assertEqual(sum(v[0] for k, v in status.items() if k != 'Total'), 10)
        self.assertEqual(status['Total'][0], 10)

        methods = block_rows(doc, 'By payment method')
        # [records, amount due, verified]
        # Every record, not only completed orders. Cash: a1 (paid) a6 b1 (paid) b3, plus
        # Cash on Delivery a5 (pending) — COD counts as cash.
        self.assertEqual(
            methods[exports.CASH],
            [5, D('1000.00') + D('700.00') + D('250.00') + D('90.00') + D('150.00'), D('1250.00')],
        )
        self.assertEqual(methods[exports.GCASH], [2, D('900.00'), D('500.00')])      # a2 (paid) + a7
        self.assertEqual(methods[exports.MAYA], [1, D('200.00'), D('0.00')])         # a4
        self.assertEqual(methods[exports.CARD], [1, D('300.00'), D('300.00')])
        self.assertEqual(methods[exports.OTHER_METHODS], [1, D('350.00'), D('0.00')])  # QR Ph b2
        self.assertEqual(sum(methods[g][0] for g in exports.PAYMENT_GROUP_ORDER), 10)
        self.assertEqual(methods['Total'][2], self.verified_amount)

    def test_detail_rows_are_the_pages_rows_and_the_totals_reconcile(self):
        client = self.client_for(self.owner)
        page_rows = client.get(reverse('billing_report'), self.range_params).context['page_obj'].paginator.object_list
        doc = self.doc('billing')
        rows = [r for g in doc.groups for r in g.rows]
        self.assertEqual(sorted(r[0] for r in rows), sorted(r.transaction.order_number for r in page_rows))
        # amount due / verified columns
        self.assertEqual(sum(r[6] for r in rows), sum(r.amount_due for r in page_rows))
        self.assertEqual(sum((r[7] or D('0')) for r in rows), self.verified_amount)
        self.assertEqual(doc.grand_total, [sum(r[6] for r in rows), self.verified_amount])
        for i in range(2):
            self.assertEqual(sum(g.subtotal[i] for g in doc.groups), doc.grand_total[i])
        self.assertNotIn(self.old.order_number, [r[0] for r in rows])      # outside the range


class InventorySummaryReconcilesTests(ExportFixture):

    def test_summary_matches_the_inventory_page_and_the_dashboard(self):
        client = self.client_for(self.owner)
        doc = self.doc('inventory')
        page = client.get(reverse('inventory_report')).context
        dash = client.get(reverse('accounts:owner_dashboard')).context
        head = headline(doc)

        self.assertEqual(head['Total products'], 5)           # five distinct products, seven stock lines
        self.assertEqual(head['Total products'], page['product_lines'])
        self.assertEqual(head['Total products'], dash['total_products_floor'])
        self.assertEqual(head['Total stock value'], page['total_value'])
        self.assertEqual(head['Total stock value'], dash['inventory_value'])
        self.assertEqual(head['Total units on hand'], page['total_units'])
        self.assertEqual(head['Total units on hand'], dash['inventory_units'])

        alerts = block_rows(doc, 'Alerts')
        # The Dashboard's Low / Critical figure is "at or below reorder level": Critical + No Stock here.
        self.assertEqual(alerts['At or below reorder level (Critical + No Stock)'][0], page['low_critical_count'])
        self.assertEqual(alerts['At or below reorder level (Critical + No Stock)'][0], dash['inventory_low_critical'])

    def test_status_counts_follow_the_stock_status_rule_and_add_up(self):
        doc = self.doc('inventory')
        status = block_rows(doc, 'Stock status')
        # Branch A: 0 / 5 / 10 / 11 / 100 against reorder 10; Branch B: 40 and 0.
        self.assertEqual(status['No Stock'][0], 2)
        self.assertEqual(status['Critical'][0], 2)            # 5 and 10 (at the reorder level)
        self.assertEqual(status['Low Stock'][0], 1)           # 11 (<= 110% of 10)
        self.assertEqual(status['In Stock'][0], 2)            # 100 and 40
        self.assertEqual(status['Total'][0], 7)
        # The same rule the Floor Stock page uses, applied to the same rows.
        expected = {}
        for inv in Inventory.objects.all():
            label = stock_status(inv.quantity_on_hand, inv.reorder_level)
            expected[label] = expected.get(label, 0) + 1
        self.assertEqual({k: v[0] for k, v in status.items() if k != 'Total'}, expected)
        self.assertEqual(status['Total'][1], headline(doc)['Total stock value'])

    def test_near_expiry_matches_the_non_performing_and_near_expiry_table(self):
        doc = self.doc('inventory')
        alerts = block_rows(doc, 'Alerts')
        page_rows = _non_performing_products(None)
        self.assertEqual(
            headline(doc)['Near-expiry items'], sum(1 for r in page_rows if r['is_near_expiry']),
        )
        self.assertEqual(headline(doc)['Near-expiry items'], 2)
        self.assertEqual(alerts['Expiring soon, within 30 days'][0], 1)
        self.assertEqual(alerts['Approaching expiry, 31-90 days'][0], 1)
        self.assertEqual(alerts['Already expired, still in stock'][0], 1)

    def test_detail_rows_subtotals_and_grand_total_reconcile(self):
        doc = self.doc('inventory')
        rows = [r for g in doc.groups for r in g.rows]
        self.assertEqual(len(rows), 7)
        self.assertEqual(sum(r[2] for r in rows), headline(doc)['Total units on hand'])
        self.assertEqual(sum(r[5] for r in rows), headline(doc)['Total stock value'])
        self.assertEqual(doc.grand_total[0], headline(doc)['Total units on hand'])
        self.assertEqual(doc.grand_total[3], headline(doc)['Total stock value'])
        self.assertEqual(sum(g.subtotal[0] for g in doc.groups), doc.grand_total[0])
        self.assertEqual(sum(g.subtotal[3] for g in doc.groups), doc.grand_total[3])
        # status column is the same rule as the Summary counts
        self.assertEqual(sorted(r[6] for r in rows).count('No Stock'), 2)


class ProcurementSummaryReconcilesTests(ExportFixture):

    def test_summary_matches_the_procurement_page_and_the_dashboard(self):
        client = self.client_for(self.owner)
        doc = self.doc('procurement')
        page = client.get(reverse('procurement_report'), self.range_params).context
        dash = client.get(reverse('accounts:owner_dashboard'), self.range_params).context
        cards = {c['key']: c for c in dash['kpi_cards']}
        head = headline(doc)
        value = block_rows(doc, 'Value')

        self.assertEqual(head['Purchase orders'], page['total_pos'])
        self.assertEqual(head['Purchase orders'], 5)
        # Ordered excludes the cancelled 9,999 and the order outside the range.
        self.assertEqual(head['Total ordered value'], D('2350.00'))
        self.assertEqual(head['Total ordered value'], page['gross_procurement'])
        self.assertEqual(head['Net procurement'], page['net_procurement'])
        self.assertEqual(cards['procurement_cost']['value_display'], f'₱{head["Net procurement"]:,.2f}')
        self.assertEqual(value['Closed POs (fully received)'][0], page['received_value'])
        self.assertEqual(value['Still open (not yet received)'][0], page['active_value'])
        returns = block_rows(doc, 'Rejected and returned')
        self.assertEqual(returns['Value of goods shipped back'][0], page['supplier_return_value'])
        self.assertEqual(returns['Supplier credit notes settled'][0], page['supplier_credit_value'])

    def test_status_table_matches_the_page_chart_and_counts_rejected_and_returned(self):
        client = self.client_for(self.owner)
        doc = self.doc('procurement')
        page = client.get(reverse('procurement_report'), self.range_params).context
        status = block_rows(doc, 'Purchase orders by status')
        self.assertEqual(
            {k: v[0] for k, v in status.items() if k != 'Total'},
            {seg['label']: seg['count'] for seg in page['chart']['segments']},
        )
        self.assertEqual(status['Total'][0], page['total_pos'])
        self.assertEqual(status['Cancelled'], [1, D('9999.00')])

        returns = block_rows(doc, 'Rejected and returned')
        self.assertEqual(returns['Deliveries rejected on receipt'][0], 1)
        self.assertEqual(returns['POs with a receipt discrepancy'][0], 1)
        self.assertEqual(returns['Supplier returns raised'][0], 1)
        self.assertEqual(returns['Units returned to suppliers'][0], 4)
        self.assertEqual(returns['Value of goods shipped back'][0], D('200.00'))

    def test_detail_rows_subtotals_and_grand_total_reconcile(self):
        doc = self.doc('procurement')
        rows = [r for g in doc.groups for r in g.rows]
        self.assertEqual(
            sorted(r[0] for r in rows),
            sorted(p.po_number for p in (self.po_received, self.po_partial, self.po_approved,
                                         self.po_cancelled, self.po_incoming)),
        )
        head = headline(doc)
        non_cancelled = [r for r in rows if r[4] != 'Cancelled']
        self.assertEqual(sum(r[7] for r in non_cancelled), head['Total ordered value'])
        self.assertEqual(sum(r[8] for r in rows), head['Total received value'])
        self.assertEqual(head['Total received value'], D('1000.00') + D('250.00'))   # 20 x 50 + 5 x 50
        self.assertEqual(sum((r[6] or 0) for r in rows), 4)
        self.assertEqual(doc.grand_total, [4, head['Total ordered value'], head['Total received value']])
        for i in range(3):
            self.assertEqual(sum(g.subtotal[i] for g in doc.groups), doc.grand_total[i])


# =============================================================================
# The PDF itself: layout, pagination, empty state
# =============================================================================

class PdfLayoutTests(ExportFixture):

    def get(self, kind, params=None, user=None, ext='pdf'):
        params = {**(self.range_params if kind != 'inventory' else {}), **(params or {})}
        return self.client_for(user or self.owner).get(reverse(f'export_{ext}', args=[kind]), params)

    def test_every_report_opens_with_header_summary_and_a_page_x_of_y_footer(self):
        titles = {
            'sales': 'Sales Report', 'billing': 'Billing Report',
            'inventory': 'Inventory Report', 'procurement': 'Procurement Report',
        }
        for kind, title in titles.items():
            with self.subTest(kind=kind):
                response = self.get(kind)
                self.assertEqual(response.status_code, 200)
                self.assertIn('attachment', response['Content-Disposition'])
                self.assertIn(f'gerlyn-{kind}-report-all-branches-', response['Content-Disposition'])
                pages = pdf_pages(response)
                first = pages[0]
                # The header's small labels are set in capitals (text-transform).
                for text in ('Gerlyn Variety Store', 'GerlynPLUS', title, 'BRANCH', 'All branches',
                             'PERIOD', 'GENERATED BY', 'Olivia Owner', 'GENERATED', 'Summary'):
                    self.assertIn(text, first)
                # The Summary is before the detail table.
                self.assertLess(first.index('Summary'), first.index('Branch: Export Branch A'))
                # Every page carries "Page k of N", resolved, with N the real page count.
                for number, text in enumerate(pages, start=1):
                    self.assertRegex(text, rf'Page {number} of {len(pages)}\b')

    def test_all_time_range_covers_the_old_rows_too(self):
        response = self.get('sales', {'range': 'all'})
        self.assertEqual(response.status_code, 200)
        text = ''.join(pdf_pages(response))
        self.assertIn('All time (through', text)
        self.assertIn(self.old.order_number, text)
        # 9,999.00 completed 40 days ago now counts: 2,750.00 + 9,999.00.
        self.assertIn('12,749.00', text)
        self.assertEqual(self.get('sales', {'range': 'all'}, ext='xlsx').status_code, 200)

    def test_filename_is_safe_whatever_the_branch_is_called(self):
        self.branch_a.branch_name = 'Baliuag "Main" ₱ — Annex'
        self.branch_a.save()
        response = self.get('sales', {'branch': self.branch_a.pk})
        self.assertEqual(response.status_code, 200)
        self.assertRegex(
            response['Content-Disposition'],
            r'^attachment; filename="gerlyn-sales-report-baliuag-main-annex-\d{4}-\d{2}-\d{2}\.pdf"$',
        )

    def test_each_report_declares_its_orientation(self):
        for kind in ('sales', 'billing', 'inventory', 'procurement'):
            with self.subTest(kind=kind):
                doc = self.doc(kind)
                box = pdf_reader(self.get(kind)).pages[0].mediabox
                if doc.orientation == 'landscape':
                    self.assertGreater(box.width, box.height)
                else:
                    self.assertGreater(box.height, box.width)

    def test_header_names_the_branch_and_period_the_report_covers(self):
        text = pdf_pages(self.get('sales', {'branch': self.branch_b.pk}))[0]
        self.assertIn('Export Branch B', text)
        self.assertNotIn('Export Branch A', text)
        self.assertIn(f'{self.start:%d %b %Y}', text)
        self.assertIn(f'{self.today:%d %b %Y}', text)
        # Inventory has no date range: it says what it is a snapshot of.
        self.assertIn('Stock on hand as of', pdf_pages(self.get('inventory'))[0])

    def test_long_table_paginates_with_headings_repeated_and_no_row_lost(self):
        # 130 more completed sales in branch A: several pages of one branch's table.
        extra = [
            SalesTransaction(
                branch=self.branch_a, total_amount=D('10.00'), transaction_type='Physical',
                mode_of_payment='Cash', order_status='Completed', recorded_by=self.owner,
            ) for _ in range(130)
        ]
        for txn in extra:
            txn.save()
        expected = {t.order_number for t in SalesTransaction.objects.exclude(pk=self.old.pk)}

        with self.assertNoLogs('xhtml2pdf', level='WARNING'):
            pages = pdf_pages(self.get('sales'))

        self.assertGreater(len(pages), 3)
        seen = []
        for number, text in enumerate(pages, start=1):
            seen += re.findall(r'ORD-\d{4}-\d+', text)
            if number > 1:
                # A continuation page repeats the branch heading and the column headings.
                self.assertIn('Order No.', text)
                self.assertRegex(text, r'Branch: Export Branch [AB]')
        # Every order exactly once: none dropped at a page break, none split across two pages.
        self.assertEqual(len(seen), len(set(seen)))
        self.assertEqual(set(seen), expected)

    def test_final_page_carries_the_grand_total_and_notes(self):
        pages = pdf_pages(self.get('sales'))
        self.assertIn('Grand total', pages[-1])
        self.assertIn('2,750.00', pages[-1])
        self.assertIn('Subtotal · Export Branch B', ''.join(pages))

    def test_no_xhtml2pdf_warnings_for_any_report(self):
        for kind in ('sales', 'billing', 'inventory', 'procurement'):
            with self.subTest(kind=kind), self.assertNoLogs('xhtml2pdf', level='WARNING'):
                self.assertEqual(self.get(kind).status_code, 200)

    def test_hostile_names_are_escaped_not_injected(self):
        self.customer.first_name, self.customer.last_name = '<b>Bold</b>', '& "Co" <script>'
        self.customer.save()
        response = self.get('sales')
        self.assertEqual(response.status_code, 200)
        self.assertIn('<b>Bold</b>', ''.join(pdf_pages(response)))     # printed literally

    def test_empty_report_prints_a_clear_no_records_page(self):
        params = {'range': 'custom', 'date_from': '2020-01-01', 'date_to': '2020-01-02'}
        for kind in ('sales', 'billing', 'procurement'):
            with self.subTest(kind=kind):
                pages = pdf_pages(self.get(kind, params))
                self.assertEqual(len(pages), 1)
                self.assertIn('No records for the selected filters', pages[0])
                self.assertIn('Gerlyn Variety Store', pages[0])
                self.assertIn('01 Jan 2020', pages[0])
                self.assertNotIn('Summary', pages[0])
                self.assertRegex(pages[0], r'Page 1 of 1\b')

    def test_empty_inventory_prints_the_no_records_page(self):
        Inventory.objects.all().delete()
        pages = pdf_pages(self.get('inventory'))
        self.assertEqual(len(pages), 1)
        self.assertIn('No records for the selected filters', pages[0])

    def test_pdf_refuses_a_report_too_large_for_it_but_excel_still_works(self):
        with mock.patch.object(exports, 'MAX_PDF_ROWS', 3):
            response = self.get('sales')
            self.assertRedirects(response, reverse('sales_report'), fetch_redirect_response=False)
            messages = [m.message for m in response.wsgi_request._messages]
            self.assertTrue(any('Export Excel' in m for m in messages), messages)
            self.assertEqual(self.get('sales', ext='xlsx').status_code, 200)


# =============================================================================
# Excel: the same Summary on its own sheet
# =============================================================================

class ExcelExportTests(ExportFixture):

    def workbook(self, kind, params=None, user=None):
        params = {**(self.range_params if kind != 'inventory' else {}), **(params or {})}
        response = self.client_for(user or self.owner).get(reverse('export_xlsx', args=[kind]), params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], exports.XLSX_CONTENT_TYPE)
        self.assertIn('.xlsx"', response['Content-Disposition'])
        return load_workbook(io.BytesIO(response.content))

    def sheet_cells(self, ws):
        return [[c.value for c in row] for row in ws.iter_rows()]

    def test_summary_sheet_comes_first_and_carries_the_header_and_headline_figures(self):
        for kind, detail_sheet in (
            ('sales', 'Sales Detail'), ('billing', 'Billing Detail'),
            ('inventory', 'Inventory Detail'), ('procurement', 'Procurement Detail'),
        ):
            with self.subTest(kind=kind):
                wb = self.workbook(kind)
                self.assertEqual(wb.sheetnames, ['Summary', detail_sheet])
                self.assertEqual(wb.active.title, 'Summary')
                doc = self.doc(kind)
                cells = self.sheet_cells(wb['Summary'])
                flat = [v for row in cells for v in row if v is not None]
                for text in (exports.STORE_NAME, exports.APP_NAME, doc.title, doc.branch_label,
                             doc.period_label, 'Olivia Owner', 'Summary', 'Key figures'):
                    self.assertIn(text, flat)
                # Every headline figure is on the sheet with its value.
                lookup = {row[0]: row[1] for row in cells if row[0] is not None}
                for label, value in doc.headline:
                    self.assertIn(label, lookup)
                    expected = float(value) if isinstance(value, Decimal) else value
                    self.assertEqual(lookup[label], expected)

    def test_summary_sheet_blocks_hold_the_same_figures_as_the_pdf_blocks(self):
        wb = self.workbook('sales')
        doc = self.doc('sales')
        cells = self.sheet_cells(wb['Summary'])
        for block in doc.summary:
            titles = [i for i, row in enumerate(cells) if row[0] == block.title]
            self.assertEqual(len(titles), 1, block.title)
            start = titles[0] + 2                                   # title, then the heading row
            for offset, expected_row in enumerate(block.rows):
                got = cells[start + offset]
                for want, have in zip(expected_row, got):
                    want = float(want) if isinstance(want, Decimal) else (want.strip() if isinstance(want, str) else want)
                    self.assertEqual(have, want)

    def test_detail_sheet_has_branch_groups_subtotals_and_grand_total_as_numbers(self):
        wb = self.workbook('sales')
        ws = wb['Sales Detail']
        cells = self.sheet_cells(ws)
        flat = [row[0] for row in cells if row[0]]
        self.assertIn('Branch: Export Branch A', flat)
        self.assertIn('Branch: Export Branch B', flat)
        self.assertTrue(any(str(v).startswith('Subtotal · Export Branch A') for v in flat))
        grand = next(row for row in cells if row[0] and str(row[0]).startswith('Grand total'))
        self.assertEqual(grand[7:], [2750.0, 100.0, 2650.0])
        # Amounts are numbers with a number format, dates are dates.
        amount = next(c for row in ws.iter_rows() for c in row if c.value == 1000.0)
        self.assertEqual(amount.number_format, '#,##0.00')
        self.assertEqual(ws.freeze_panes, 'A5')
        self.assertEqual(ws.page_setup.orientation, 'landscape')

    def test_detail_rows_match_the_pdf_rows(self):
        wb = self.workbook('billing')
        doc = self.doc('billing')
        order_numbers = {row[0] for row in self.sheet_cells(wb['Billing Detail'])
                         if row[0] and str(row[0]).startswith('ORD-')}
        self.assertEqual(order_numbers, {r[0] for g in doc.groups for r in g.rows})

    def test_a_formula_in_a_name_is_stored_as_text_not_executed(self):
        self.customer.first_name, self.customer.last_name = "=HYPERLINK(\"http://x\")", '+cmd'
        self.customer.save()
        ws = self.workbook('sales')['Sales Detail']
        cell = next(c for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith('=HYPERLINK'))
        self.assertEqual(cell.data_type, 's')

    def test_empty_report_says_so_on_both_sheets(self):
        params = {'range': 'custom', 'date_from': '2020-01-01', 'date_to': '2020-01-02'}
        wb = self.workbook('sales', params)
        for name in ('Summary', 'Sales Detail'):
            flat = [v for row in self.sheet_cells(wb[name]) for v in row if v]
            self.assertIn('No records for the selected filters', flat)

    def test_inventory_excel_reflects_the_branch_pick(self):
        ws = self.workbook('inventory', {'branch': self.branch_b.pk})['Inventory Detail']
        flat = [row[0] for row in self.sheet_cells(ws) if row[0]]
        self.assertIn('Branch: Export Branch B', flat)
        self.assertNotIn('Branch: Export Branch A', flat)


# =============================================================================
# Who may export what
# =============================================================================

class ExportAccessTests(ExportFixture):

    def url(self, kind='sales', ext='pdf'):
        return reverse(f'export_{ext}', args=[kind])

    def test_anonymous_is_sent_to_login(self):
        for ext in ('pdf', 'xlsx'):
            response = Client().get(self.url(ext=ext), self.range_params)
            self.assertEqual(response.status_code, 302)
            self.assertIn('login', response['Location'])

    def test_staff_and_customers_cannot_export(self):
        staff = User.objects.create_user(username='exp-staff', password='pw', role=User.Role.STAFF)
        for user in (staff, self.customer, self.supplier):
            for ext in ('pdf', 'xlsx'):
                with self.subTest(user=user.username, ext=ext):
                    response = self.client_for(user).get(self.url(ext=ext), self.range_params)
                    self.assertEqual(response.status_code, 403)

    def test_manager_gets_only_their_branch_whatever_the_query_says(self):
        for params in ({}, {'branch': self.branch_b.pk}):
            with self.subTest(params=params):
                doc = self.doc('sales', user=self.manager, **({'branch_id': self.branch_b.pk} if params else {}))
                self.assertEqual(doc.branch_label, 'Export Branch A')
                self.assertEqual([g.title for g in doc.groups], ['Export Branch A'])
                self.assertEqual(headline(doc)['Gross sales'], D('2150.00'))
                response = self.client_for(self.manager).get(self.url(), {**self.range_params, **params})
                text = ''.join(pdf_pages(response))
                self.assertIn('Export Branch A', text)
                self.assertNotIn('Export Branch B', text)

    def test_unassigned_manager_fails_closed_and_the_header_says_why(self):
        orphan = User.objects.create_user(username='exp-orphan', password='pw', role=User.Role.BRANCH_MANAGER)
        for kind in ('sales', 'billing', 'inventory', 'procurement'):
            with self.subTest(kind=kind):
                doc = self.doc(kind, user=orphan)
                self.assertTrue(doc.is_empty)
                self.assertEqual(doc.branch_label, 'No branch assigned')

    def test_bad_range_redirects_back_to_the_report_with_a_message(self):
        for ext in ('pdf', 'xlsx'):
            response = self.client_for(self.owner).get(self.url(ext=ext), {'range': 'custom'})
            self.assertRedirects(response, reverse('sales_report'), fetch_redirect_response=False)

    def test_unknown_report_is_a_404_and_a_junk_branch_is_ignored(self):
        self.assertEqual(self.client_for(self.owner).get(self.url('payroll')).status_code, 404)
        response = self.client_for(self.owner).get(self.url(), {**self.range_params, 'branch': 'abc'})
        self.assertEqual(response.status_code, 200)

    def test_report_pages_link_to_the_new_exports(self):
        client = self.client_for(self.owner)
        pages = {
            'sales': ['reports_dashboard', 'sales_report'], 'billing': ['billing_report'],
            'inventory': ['inventory_report'], 'procurement': ['procurement_report'],
        }
        for kind, names in pages.items():
            for name in names:
                with self.subTest(page=name):
                    html = client.get(reverse(name), self.range_params).content.decode()
                    self.assertIn(reverse('export_pdf', args=[kind]), html)
                    self.assertIn(reverse('export_xlsx', args=[kind]), html)
                    self.assertNotIn('onclick="window.print()"', html)

    def test_export_link_carries_the_pages_branch_and_range(self):
        html = self.client_for(self.owner).get(
            reverse('billing_report'), {**self.range_params, 'branch': self.branch_a.pk},
        ).content.decode()
        link = re.search(r'href="(/reports/export/billing\.pdf\?[^"]*)"', html).group(1)
        self.assertIn(f'branch={self.branch_a.pk}', link)
        self.assertIn('range=custom', link)
        # ...and following it describes that branch only.
        text = ''.join(pdf_pages(self.client_for(self.owner).get(link.replace('&amp;', '&'))))
        self.assertIn('Export Branch A', text)
        self.assertNotIn('Branch: Export Branch B', text)
